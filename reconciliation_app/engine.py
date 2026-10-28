"""
RECONCILIATION ENGINE
=====================
The rules that are the SAME FOR EVERY BANK. The app (app.py) calls these functions; the developer notebook
(I_M_reconciliation.ipynb) uses the review functions of section 7.

How the code is split
    app.py           the screens (6 steps + side panel); no reconciliation rules
    engine.py        this file: the common rules
    banks/<bank>.py  what is specific to one bank: statement format, matching, SLA, calculations
                     (banks/__init__.py lists the banks and explains how to add one)

What happens in one reconciliation (the order of this file)
    1. choose the bank and the period (day / week / month) ............ sections 1, 3
    2. read the uploaded files in memory (statements + IremboPay file) . section 5
    3. check that every day is covered before Run ...................... section 6
    4. run: match the bank rows with IremboPay, review the rows that are
       not found, check the SLA, build the 5 summary tables ............ sections 7, 9
    5. compare with the previous report (optional) .................... section 8
    6. write the Excel report, log the run, make the PDF ............... sections 10, 11, 13

What is saved (nothing else: uploaded files are never kept)
    output/<day|week|month>/<CODE>_recon_summary_<period>.xlsx   the reports  (CODE = IM, ECO, ...)
    output/.review_store.csv                                       the review decisions (shared with the notebook)
    output/run_log.csv                                             one line per run

CONTENTS
    1.  Errors and banks
    2.  Settings
    3.  Periods
    4.  File names and columns
    5.  Reading the uploaded files
    6.  Checks before Run
    7.  Review decisions
    8.  Comparison with the previous report
    9.  The reconciliation run
    10. Run log
    11. Writing the Excel report
    12. Saved reports and decisions (side panel)
    13. PDF (summary only)
    14. Custom report (step 6: sections, charts and layout chosen by the user)

Words used in the code
    statement row   one line of a bank statement (one payment, or for Ecobank one payment attempt)
    payment date    the bank's date of the payment (I&M: Value Date, Ecobank: Request Date); sets the period
    not found       a statement row whose invoice is not in the IremboPay file: it must be reviewed
    block           one summary of 5 tables (a day, a week or a month); a monthly report has one block per week
                    + one for the month
"""
from __future__ import annotations

import hashlib
import io
import re
import time
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import banks


# =====================================================================================================================
# 1. ERRORS AND BANKS
#    ReconError = a problem shown to the user in plain words. The banks come from banks/ (see banks/__init__.py).
# =====================================================================================================================
class ReconError(Exception):
    """A problem explained in plain words for the user."""


class NoTransactionsError(ReconError):
    """The files cover the period, but none of their rows fall in it: the bank had no transactions. Not really a
    problem - the app answers it with a No-Transaction Report (no_transaction_excel) instead of showing an error
    (approved 2026-10-27)."""


BANKS = banks.load()                       # {code: bank module}; raises ValueError if a bank file is incomplete
DEFAULT_BANK = banks.ACTIVE[0].CODE        # I&M: used when no bank is given (e.g. by the notebook)


def get_bank(bank=None):
    """A bank module from its code (or the module itself); None = the default bank (I&M)."""
    if bank is None:
        return BANKS[DEFAULT_BANK]
    if not isinstance(bank, str):
        return bank
    if bank not in BANKS:
        raise ReconError(f"Unknown bank '{bank}'.")
    return BANKS[bank]


def bank_name(code: str) -> str:
    """'IM' -> 'I&M' (an unknown code is returned as it is)."""
    return BANKS[code].NAME if code in BANKS else code


def option(bank, name, default):
    """Optional bank setting (see banks/__init__.py OPTIONAL); the default is how I&M works."""
    return getattr(get_bank(bank), name, default)


def present(*cols):
    """Columns that the bank has (a bank may have no posted date)."""
    return [c for c in cols if c]


# =====================================================================================================================
# 2. SETTINGS
#    Column names of the IremboPay file, the review choices, the report sheets and file names.
# =====================================================================================================================
# ---- IremboPay payment file
TEXT_COLUMNS = ["invoice_number", "transaction_id", "provider_reference", "transaction_reference"]  # kept as text
DATE_COLUMNS = ["created_at", "paid_at", "settled_at", "expiry_at"]                                  # read as dates
PAYMENT_REQUIRED = ["invoice_number", "amount", "currency", "payment_status", "paid_at",
                    "payment_provider", "transaction_status", "created_at", "name"]                  # must be there

# ---- review (step 4 of the app): the 5 choices and the columns the user fills in
TABLE_CHOICES = ["Notification received", "Pending notification", "Paid with other channel",
                 "Failed notification", "Exclude"]
# 26/10/2026: an extra credit / debit (a second one for an invoice already counted) is not in Processed, but the
# table chosen for it in Exception Resolution is read: it counts there (a second payment = a double payment)
EXTRA_DECIDED = {"Notification received": "received", "Pending notification": "pending",
                 "Paid with other channel": "duplicate"}


EXCLUDED = "Excluded (not counted)"          # 26/10/2026: an excluded row stays in the report, with this status
NOT_COUNTED = "Not counted (second credit)"   # an extra credit in no table (Failed notification / no decision)


def with_extra_decisions(rows: pd.DataFrame, m: dict, window: pd.Series) -> dict:
    """The tables `m` (masks) plus the extra credits of the window put in a table in Exception Resolution
    (Failed notification / Exclude / no decision: in no table). A row excluded in the review is in no table
    (26/10/2026: it is not removed from the report any more - it stays in the statement, with a comment)."""
    if "review_table" not in rows:
        return m
    table = rows["review_table"].fillna("").astype(str)
    excluded = table.eq("Exclude")
    out = {k: v & ~excluded for k, v in m.items()}
    if "is_extra_credit" in rows:
        # 26/10/2026: an invoice counts at most ONCE in each table - an extra credit put in a table that already has
        # its invoice (its counted credit, or an earlier extra one) is not counted again (it stays, with a note)
        extra = window & rows["is_extra_credit"].fillna(False).astype(bool) & ~excluded
        inv = rows["invoice_number"].astype(object).where(rows["invoice_number"].notna(), "").astype(str)
        for choice, key in EXTRA_DECIDED.items():
            there = set(inv[out[key]])
            cand = inv[extra & table.eq(choice) & inv.ne("") & ~inv.isin(there)]
            add = pd.Series(False, index=rows.index)
            add[cand[~cand.duplicated()].index] = True       # the first of them only
            out[key] = out[key] | add
    return out


TABLE_NAMES = {"received": "Notification received", "pending": "Pending notification",
               "duplicate": "Duplicate payment"}


def table_notes(rows: pd.DataFrame, whole: dict, bank) -> pd.Series:
    """26/10/2026: "Report note" of each row of an invoice credited / debited more than once - why the row is in its
    table (or in none) and where the other credit of the invoice went."""
    bank = get_bank(bank)
    ref, day = bank.COLUMNS.get("reference"), bank.COLUMNS["payment_date"]
    notes = pd.Series("", index=rows.index, dtype=object)
    if "is_extra_credit" not in rows:
        return notes
    inv = rows["invoice_number"].astype(object).where(rows["invoice_number"].notna(), "").astype(str)
    many = inv.ne("") & inv.duplicated(keep=False)
    where = pd.Series("", index=rows.index, dtype=object)
    for key, name in TABLE_NAMES.items():
        where[whole[key] & where.eq("")] = name
    extra = rows["is_extra_credit"].fillna(False).astype(bool)
    table = rows["review_table"].fillna("").astype(str) if "review_table" in rows else pd.Series("", index=rows.index)

    def shown(choice):                       # the table a review choice puts a row in
        return TABLE_NAMES[EXTRA_DECIDED[choice]]

    def label(i):
        r = str(rows.at[i, ref]) if ref and ref in rows and pd.notna(rows.at[i, ref]) else ""
        d = pd.to_datetime(rows.at[i, day], errors="coerce")
        return " ".join(x for x in (r, f"of {d:%d/%m %H:%M}" if pd.notna(d) else "") if x)

    for _, idx in rows[many].groupby(inv[many]).groups.items():
        for i in idx:
            if where[i]:
                why = (f"In {where[i]}: " + ("second credit of this invoice, put here in Exception Resolution"
                                             if extra[i] else "the credit counted for this invoice"))
            elif extra[i] and table[i] in EXTRA_DECIDED:
                why = f"Not counted: the invoice is already in {shown(table[i])}"
            elif extra[i]:
                why = "Not counted: second credit of this invoice" + (" (excluded)" if table[i] == "Exclude" else "")
            else:
                why = "The credit counted for this invoice"
            others = [f"{label(j)} -> {where[j] or (f'not counted (already in {shown(table[j])})' if extra[j] and table[j] in EXTRA_DECIDED else 'no table')}"
                      for j in idx if j != i]
            notes[i] = why + ". Other credit(s) of the invoice: " + "; ".join(others) + "."
    return notes


EDIT_COLS = ["Reviewed (Yes/No)", "Report table", "SLA violated (Yes/No)", "Correct invoice number",
             "payment_status", "payment_provider", "paid_at (YYYY-MM-DD HH:MM:SS)", "Merchant name", "Comment",
             "Currency (chosen)"]                  # the last one: banks whose statement has no currency (BPR)
CURRENCY_CHOSEN = "Currency (chosen)"
SUGGEST_COL = "Suggestion from"            # filled when an earlier review already decided the invoice
STORE_COLS = ["bank", "period", "key", "invoice_number"] + EDIT_COLS + ["saved_at"]   # columns of the review store
STORE_NAME = ".review_store.csv"

# ---- report: the 5 summary tables, the sheets, the file names
TABLE_TITLES = ["Overall processed payments", "Overall notifications received by Irembo", "SLA Violated",
                "Pending notifications", "Duplicate Payment"]
SUMMARY_SHEET = {"day": "Summary-Daily", "week": "Summary-Weekly", "month": "Summary-Monthly"}
# the 6th table of each summary block (2026-10-12): the notifications received, by merchant and currency
MERCHANT_TITLE = "Payment per merchant"
OLD_SHEETS = {"Payment per merchant", "Debited more than once"}   # Ecobank sheets of reports made before: removed
NOT_FOUND_SHEET = "not found in IremboPay"     # only in old reports (removed when they are made again)
REPORT_SHEETS = ({"Summary-Daily", "Summary-Weekly", "Summary-Monthly", "Pending notification",
                  "Duplicate payment other channel", NOT_FOUND_SHEET, "Payment details"}
                 | {b.STATEMENT_SHEET for b in BANKS.values()}
                 | {s for b in BANKS.values() for s in getattr(b, "EXTRA_SHEETS", [])})   # written by the app (rest = user's)
REPORT_SUFFIX = "_recon_summary_"              # report file = <CODE>_recon_summary_<period label>.xlsx

# ---- fast Excel reading when python-calamine is installed (Setup installs it); pandas' default otherwise
try:
    import python_calamine  # noqa: F401
    EXCEL_ENGINE = "calamine"
except ImportError:
    EXCEL_ENGINE = None


# =====================================================================================================================
# 3. PERIODS
#    A period is a day, a week (Monday to Sunday) or a month; `end` is the midnight after the last day (NOT included).
#    Only a period that has ended can be reconciled (period_not_ended).
# =====================================================================================================================
@dataclass
class Period:
    report_type: str           # "day", "week" or "month"
    start: pd.Timestamp        # first day
    end: pd.Timestamp          # not included
    unit: str                  # "Day", "Week" or "Month" (used in "VS prev <unit>")

    @property
    def label(self):
        """'01-09-2026_to_30-09-2026' (first day to last day): used in report file names."""
        return f"{self.start:%d-%m-%Y}_to_{self.last_day:%d-%m-%Y}"

    @property
    def store_key(self):
        """'month:01-09-2026_to_01-10-2026': the period of a review decision. It keeps the day AFTER the period as
        its end: the format of the decisions already saved."""
        return f"{self.report_type}:{self.start:%d-%m-%Y}_to_{self.end:%d-%m-%Y}"

    @property
    def text(self):
        """'01/09/2026 - 30/09/2026': used in the run log."""
        return f"{self.start:%d/%m/%Y} - {self.last_day:%d/%m/%Y}"

    @property
    def last_day(self):
        return self.end - pd.Timedelta(days=1)

    @property
    def name(self):
        """How the period is shown to the user: 'September 2026', 'week 07/09 - 13/09/2026', ..."""
        return {"day": f"{self.start:%A %d %B %Y}",
                "week": "w" + f"{week_name(self)} ({self.start:%d/%m} - {self.last_day:%d/%m/%Y})"[1:],
                "month": f"{self.start:%B %Y}"}[self.report_type]

    @property
    def day_before(self):
        return self.start - pd.Timedelta(days=1)

    @property
    def monday_after(self):
        """Monday after the period when it ends on a weekend (late weekend credits are posted on Monday)."""
        if self.last_day.dayofweek >= 5:
            return self.last_day + pd.Timedelta(days=7 - self.last_day.dayofweek)
        return None

    def previous(self) -> "Period":
        """The period just before (used by the comparison)."""
        return make_period(self.report_type, {"day": self.start - pd.Timedelta(days=1),
                                              "week": self.start - pd.Timedelta(days=7),
                                              "month": self.start - pd.offsets.MonthBegin(1)}[self.report_type])


def month_of(kind: str, start: pd.Timestamp) -> str:
    """'2026-10': the month a period starting on `start` belongs to - a week by its Thursday, a day or a month by
    its own start (kind: "day" / "week" / "month", as on Period.report_type or a saved report's name)."""
    anchor = start + pd.Timedelta(days=3) if kind == "week" else start
    return f"{anchor:%Y-%m}"


def make_period(report_type: str, date) -> Period:
    """The day / week / month that contains `date`."""
    day = pd.Timestamp(date).normalize()
    if report_type == "day":
        return Period("day", day, day + pd.Timedelta(days=1), "Day")
    if report_type == "week":
        start = day - pd.Timedelta(days=day.dayofweek)          # back to Monday
        return Period("week", start, start + pd.Timedelta(days=7), "Week")
    if report_type == "month":
        start = day.replace(day=1)
        return Period("month", start, start + pd.offsets.MonthBegin(1), "Month")
    raise ReconError('Report type must be "day", "week" or "month".')


def period_not_ended(period: Period, today=None) -> str | None:
    """Why a period cannot be reconciled yet, or None when it has ended. A period ends at midnight after its last
    day: a day the next day, a week (Monday to Sunday) on the next Monday, a month on the 1st of the next month."""
    today = pd.Timestamp(today if today is not None else pd.Timestamp.today()).normalize()
    if period.end <= today:
        return None
    what = {"day": f"The day {period.start:%A %d/%m/%Y}",
            "week": f"The week {period.start:%d/%m/%Y} – {period.last_day:%d/%m/%Y}",
            "month": f"The month {period.start:%B %Y}"}[period.report_type]
    return f"{what} has not ended yet.\nIt can be reconciled from {period.end:%A %d/%m/%Y}."


# ---- what IremboPay says about a matched row (26/10/2026, every bank): paid = paid_at is filled, whatever the
#      payment_status; the decision of Exception Resolution always wins
def is_paid(rows: pd.DataFrame) -> pd.Series:
    table = rows["review_table"].fillna("") if "review_table" in rows else pd.Series("", index=rows.index)
    paid = rows["paid_at"].notna() | table.isin(["Notification received", "Paid with other channel"])
    return paid & table.ne("Pending notification")


def bank_providers(bank) -> set:
    """The IremboPay payment_provider values that mean this bank: PROVIDER, plus PROVIDER_ALIASES (27/10/2026, BPR:
    'BPR' - a bank transfer into IremboPay's BPR account - is BPR too, like 'BPR_RW')."""
    bank = get_bank(bank)
    return {bank.PROVIDER, *option(bank, "PROVIDER_ALIASES", [])}


def _providers(provider) -> set:
    return set(provider) if isinstance(provider, (set, list, tuple, frozenset)) else {provider}


def paid_through(rows: pd.DataFrame, provider) -> pd.Series:
    """Notification received: in IremboPay, paid (paid_at), through this bank (provider: one name or several)."""
    return (rows["found_in_irembopay"].fillna(False).astype(bool) & rows["payment_provider"].isin(_providers(provider))
            & is_paid(rows))


def paid_elsewhere(rows: pd.DataFrame, provider) -> pd.Series:
    """Duplicate Payment: in IremboPay, paid (paid_at), but through another provider (not one of this bank's)."""
    return (rows["found_in_irembopay"].fillna(False).astype(bool) & rows["payment_provider"].notna()
            & ~rows["payment_provider"].isin(_providers(provider)) & is_paid(rows))


def not_paid(rows: pd.DataFrame, pending_statuses) -> pd.Series:
    """Pending notification: in IremboPay but not paid (no paid_at), or a row the bank's rules make pending (BPR: a
    credit with no invoice number, status NEW)."""
    known = rows["found_in_irembopay"].fillna(False).astype(bool) | rows["payment_status"].isin(pending_statuses)
    return known & ~is_paid(rows)


def week_number(monday, month=None) -> tuple:
    """(month, n) of the Monday-to-Sunday week starting `monday` (decided 23/10/2026). Its month: the month given
    (a block of a monthly report), else the month where most of its days are (the month of its Thursday). n = its
    place among the Monday-to-Sunday weeks that touch that month: the week of the 1st is Week 1 (also when it starts
    in the month before). E.g. 28/09-04/10/2026 = October Week 1; 07/09-13/09/2026 = September Week 2."""
    monday = pd.Timestamp(monday).normalize()
    if month is None:
        month = (monday + pd.Timedelta(days=3)).replace(day=1)
    first = pd.Timestamp(month).replace(day=1)
    base = first - pd.Timedelta(days=first.dayofweek)
    return first, (monday - base).days // 7 + 1


def week_name(period: "Period", short: bool = False) -> str:
    """'Week 1 · October 2026' (short: 'Week 1') of a weekly period; '' for a day or a month."""
    if period.report_type != "week":
        return ""
    month, n = week_number(period.start)
    return f"Week {n}" if short else f"Week {n} · {month:%B %Y}"


# public holidays (Settings page, app_settings.json): a payment made on a weekend or a holiday is posted on the next
# working day. Set by the app at each run (set_holidays); empty = weekends only.
HOLIDAYS: dict = {}                       # Timestamp -> name


def set_holidays(holidays) -> None:
    """holidays: [{"date": "2026-04-06", "name": "Easter Monday"}, ...] (or (date, name) pairs)."""
    HOLIDAYS.clear()
    for h in holidays or []:
        date, name = (h.get("date"), h.get("name", "")) if isinstance(h, dict) else h
        day = pd.to_datetime(date, errors="coerce")
        if pd.notna(day):
            HOLIDAYS[pd.Timestamp(day).normalize()] = str(name or "Holiday")


def is_working_day(day) -> bool:
    day = pd.Timestamp(day).normalize()
    return day.dayofweek < 5 and day not in HOLIDAYS


def next_working_day_after(period: Period):
    """When the period ends on a weekend or a holiday: the first working day after it (its payments are posted
    then), skipping weekends and holidays (e.g. a week ending Sunday 05/04/2026 -> Wednesday 08/04: Easter Monday
    and 07/04 are holidays). None when the period ends on a working day."""
    day = period.last_day
    if is_working_day(day):
        return None
    day = day + pd.Timedelta(days=1)
    while not is_working_day(day):
        day = day + pd.Timedelta(days=1)
    return day


def extra_monday(period: Period, bank=None):
    """The next working day after the period, when the bank needs it (MONDAY_AFTER_WEEKEND: the period ends on a
    weekend or a holiday, whose payments are posted on the next working day), else None. Usually a Monday."""
    return next_working_day_after(period) if get_bank(bank).MONDAY_AFTER_WEEKEND else None


def closed_days_before(day) -> list:
    """The weekend days and holidays just before a working day (the days whose payments it holds), oldest first."""
    out, d = [], pd.Timestamp(day).normalize() - pd.Timedelta(days=1)
    while not is_working_day(d):
        out.insert(0, d)
        d -= pd.Timedelta(days=1)
    return out


# =====================================================================================================================
# 4. FILE NAMES AND COLUMNS
# =====================================================================================================================
def report_path(base_dir: Path, period: Period, bank=None, root: Path | None = None) -> Path:
    """Where the report of a bank and period is written: <root>/<day|week|month>/<CODE>_recon_summary_<label>.xlsx.
    root: output/ by default; the app gives its session's temporary folder (reports are kept on Google Drive)."""
    code = get_bank(bank).CODE
    root = Path(root) if root is not None else Path(base_dir) / "output"
    week = f" ({week_name(period, short=True).lower()})" if period.report_type == "week" else ""   # ' (week 1)'
    return root / period.report_type / f"{code}{REPORT_SUFFIX}{period.label}{week}.xlsx"


def report_bank(path: Path) -> str:
    """Bank code from a report file name ('IM_recon_summary_...' -> 'IM')."""
    return Path(path).stem.split(REPORT_SUFFIX)[0]


def report_label(path: Path) -> str:
    """Period label from a report file name ('IM_recon_summary_01-09-2026_to_30-09-2026' -> '01-09-2026_to_...')."""
    return re.sub(r" \(week \d+\)", "", Path(path).stem.split(REPORT_SUFFIX, 1)[-1])   # without ' (week 1)'


def bank_cols(bank=None) -> list:
    """Columns of a bank row shown in the review (not found in IremboPay)."""
    c = get_bank(bank).COLUMNS
    return (["Currency"] + present(c["payment_date"], c["posted_date"], c["description"], c["reference"], c["amount"])
            + ["invoice_number", "Bank channel", "Reason", "Source file"])


# =====================================================================================================================
# 5. READING THE UPLOADED FILES (in memory, never saved)
#    A bank statement is read by the bank's own read_statement(); the IremboPay file is the same for every bank.
# =====================================================================================================================
@dataclass
class UploadedFile:
    name: str
    kind: str                  # "bank" or "payments"
    digest: str                # fingerprint of the content: the same file added twice is recognised
    data: pd.DataFrame
    currency: str = ""
    first: pd.Timestamp | None = None      # first day covered
    last: pd.Timestamp | None = None       # last day covered
    rows: int = 0
    other_credits: int = 0                 # credits on the statement that are not IremboPay payments
    bank: str = ""                         # bank code (bank files only)
    providers: frozenset | None = None     # payment_provider values found (payment files only, e.g. {"I_M"})
    balances: dict | None = None           # opening / closing / net of a statement (BPR: the checks between files)


def file_digest(content: bytes) -> str:
    """Fingerprint of a file's content (two files with the same content have the same fingerprint)."""
    return hashlib.sha1(content).hexdigest()


def read_bank_file(name: str, content: bytes, bank=None) -> UploadedFile:
    """Read one bank statement with the bank's rules; the days it covers come from its header dates, else from
    the payment dates inside."""
    bank = get_bank(bank)
    try:
        data, other_credits, info = bank.read_statement(content, name, EXCEL_ENGINE)
    except ValueError as err:                                     # the bank file explains what is wrong
        raise ReconError(str(err)) from None
    data["Source file"] = name
    first = info["date_from"] if pd.notna(info["date_from"]) else info["first"]
    last = info["date_to"] if pd.notna(info["date_to"]) else info["last"]
    if pd.isna(first) or pd.isna(last):
        raise ReconError(f"{name}: no dates found in this bank file.")
    return UploadedFile(name, "bank", file_digest(content), data, info["currency"],
                        pd.Timestamp(first).normalize(), pd.Timestamp(last).normalize(), len(data), other_credits,
                        bank.CODE, balances=info.get("balances"))


def read_payment_file(name: str, content: bytes, bank=None) -> UploadedFile:
    """Read the IremboPay payment file (.csv, or .xlsx when the bank allows it); the days it covers come from
    paid_at."""
    is_excel = name.lower().endswith((".xlsx", ".xls")) and "xlsx" in option(bank, "PAYMENT_FILE_TYPES", ["csv"])
    try:
        if is_excel:
            df = pd.read_excel(io.BytesIO(content), dtype={c: str for c in TEXT_COLUMNS}, engine=EXCEL_ENGINE)
        else:
            df = pd.read_csv(io.BytesIO(content), dtype={c: str for c in TEXT_COLUMNS})
    except Exception:
        kind = "an Excel" if is_excel else "a CSV"
        raise ReconError(f"{name}: this file cannot be read as {kind} payment file.") from None
    optional = option(bank, "PAYMENT_OPTIONAL_COLUMNS", [])
    missing = [c for c in PAYMENT_REQUIRED if c not in df.columns and c not in optional]
    if missing:
        raise ReconError(f"{name}: this is not an IremboPay payment file (missing {', '.join(missing)}).")
    for c in DATE_COLUMNS:
        if c in df:
            df[c] = pd.to_datetime(df[c], errors="coerce")
    df["amount"] = pd.to_numeric(df["amount"], errors="coerce")
    df["invoice_number"] = df["invoice_number"].astype(str).str.strip()
    if "created_at" not in df:                      # optional for some banks: paid_at is used instead
        df["created_at"] = df["paid_at"]
    paid = df["paid_at"].dropna()
    if paid.empty:
        raise ReconError(f"{name}: no paid_at dates found in this payment file.")
    providers = frozenset(df["payment_provider"].dropna().astype(str).str.strip().str.upper())
    return UploadedFile(name, "payments", file_digest(content), df, "",
                        paid.min().normalize(), paid.max().normalize(), len(df), providers=providers)


# =====================================================================================================================
# 6. CHECKS BEFORE RUN (step 3 of the app)
#    accept_file: is an uploaded file usable for this period?   check_period: is every day covered?
# =====================================================================================================================
@dataclass
class Check:
    blocking: list = field(default_factory=list)     # Run is blocked
    warnings: list = field(default_factory=list)
    lines: list = field(default_factory=list)        # (ok: bool, text) shown as a checklist
    missing: list = field(default_factory=list)      # days with no statement that do not block, e.g. "17/10 to 18/10"


def _ranges(days):
    """[1/9, 2/9, 3/9, 7/9] -> [(1/9, 3/9), (7/9, 7/9)]: missing days shown as ranges."""
    days, out = sorted(days), []
    for d in days:
        if out and (d - out[-1][1]).days == 1:
            out[-1][1] = d
        else:
            out.append([d, d])
    return [(a, b) for a, b in out]


def accept_file(period: Period, f: UploadedFile, accepted: list, bank=None) -> str | None:
    """Reason to reject an uploaded file, or None if it is accepted."""
    bank = get_bank(bank)
    if any(a.digest == f.digest for a in accepted):
        return f"{f.name}: the same file was already added."
    if f.kind == "bank" and f.bank != bank.CODE:
        return f"{f.name}: this is a {bank_name(f.bank)} statement, not {bank.NAME}."
    if f.kind == "payments" and f.providers is not None and bank.PROVIDER.upper() not in f.providers:
        return f"{f.name}: this IremboPay file has no {bank.NAME} payments (payment_provider {bank.PROVIDER})."
    # the days a file may cover: statements from the day before (late payments) to the Monday after a weekend
    if f.kind == "bank":
        window_start, window_end = period.day_before, extra_monday(period, bank) or period.last_day
        if option(bank, "DAY_AFTER_RECOMMENDED", False):
            window_end = max(window_end, period.last_day + pd.Timedelta(days=1))
        earlier = option(bank, "ACCEPT_EARLIER_DAYS", 0)      # payments of the period requested before it
        if earlier:
            window_start = period.start - pd.Timedelta(days=earlier)
    else:
        window_start, window_end = period.start, period.last_day
    if f.last < window_start or f.first > window_end:
        return (f"{f.name}: this file covers {f.first:%d/%m/%Y} to {f.last:%d/%m/%Y}, outside "
                f"{period.name} ({period.start:%d/%m/%Y} to {period.last_day:%d/%m/%Y}).")
    return None


def check_period(period: Period, files: list, bank=None, today=None) -> Check:
    """The checklist of step 3: which days are covered by statements, what is missing, what is recommended.
    today: None = the real today (another day only for testing)."""
    bank = get_bank(bank)
    chk = Check()
    today = pd.Timestamp(today if today is not None else pd.Timestamp.today()).normalize()
    statements = [f for f in files if f.kind == "bank"]
    pays = [f for f in files if f.kind == "payments"]
    not_ended = period_not_ended(period, today)
    if not_ended:
        chk.lines.append((False, not_ended.replace("\n", " ")))
        chk.blocking.append(not_ended)
    monday = extra_monday(period, bank)
    needed = list(pd.date_range(period.start, period.last_day))    # every day: only ended periods are reconciled
    if monday is not None:                                         # always needed, even when not over yet
        needed.append(monday)

    # ---- statements: every needed day must be covered (per currency, or once when one file holds all currencies)
    per_currency = option(bank, "ONE_FILE_PER_CURRENCY", True)
    block_missing = option(bank, "MISSING_DAYS_BLOCK", True)
    groups = ([(cur, f"Bank {cur}", f"the {cur} statement", [f for f in statements if f.currency == cur])
               for cur in bank.CURRENCIES] if per_currency
              else [(None, f"{bank.NAME} statements", f"the {bank.NAME} statement", statements)])
    for cur, label, what, mine in groups:
        covered = set()
        for f in mine:
            covered |= set(pd.date_range(f.first, f.last))
        if hasattr(bank, "check_files") and mine:          # the balances: nothing missing between the files
            covered, more_lines, more_blocking = bank.check_files(mine)
            chk.lines += more_lines
            chk.blocking += more_blocking
        if not mine:
            chk.lines.append((False, f"{label}: no file added → upload {what}(s) for "
                                     f"{period.start:%d/%m} to {(monday or period.last_day):%d/%m/%Y}"))
            chk.blocking.append(f"{label}: no file added.")
        else:
            gaps = [d for d in needed if d not in covered]
            for a, b in _ranges(gaps):
                txt = f"{a:%A %d/%m}" if a == b else f"{a:%d/%m} to {b:%d/%m}"
                why = ""
                if monday is not None and a <= monday <= b:
                    closed = closed_days_before(monday)
                    names = [HOLIDAYS[d] for d in closed if d in HOLIDAYS]
                    why = (f" (next working day: payments of {closed[0]:%d/%m}" + (f" to {closed[-1]:%d/%m}"
                           if len(closed) > 1 else "") + (f", {' and '.join(names)}," if names else "")
                           + f" are posted on {monday:%A %d/%m}")
                    why += (f"; its statement is available from {monday + pd.Timedelta(days=1):%A %d/%m})"
                            if monday >= today else ")")
                chk.lines.append((False, f"{label}: {txt} missing → upload {what} for "
                                         f"{'that day' if a == b else 'those days'}{why}"))
            if gaps and block_missing:
                chk.blocking.append(f"{label}: some days are missing.")
            elif gaps:                                         # some banks: a warning only, Run stays possible
                days_txt = ", ".join(f"{a:%d/%m}" if a == b else f"{a:%d/%m} to {b:%d/%m}" for a, b in _ranges(gaps))
                chk.missing.append(f"{label}: {days_txt}")
                chk.warnings.append(f"{label}: no statement for {days_txt}. Run is allowed: those days may simply have "
                                    "had no transaction. If they had some, upload their statement, otherwise their "
                                    "payments will be missing from the report. The Slack message will say it.")
            else:
                chk.lines.append((True, f"{label}: {period.start:%d/%m} to {max(needed):%d/%m} covered"))
        # the statement of the day before may hold payments of the first day (recommended, never blocking)
        day_after = period.last_day + pd.Timedelta(days=1)
        if (option(bank, "DAY_AFTER_RECOMMENDED", False) and day_after < today and day_after not in covered
                and mine):
            chk.warnings.append(f"Recommended: upload the statement of {day_after:%d/%m/%Y}. Payments made late on "
                                f"{period.last_day:%d/%m/%Y} are booked the next morning.")
        if cur in bank.DAY_BEFORE_CURRENCIES and period.day_before < today and period.day_before not in covered:
            chk.warnings.append(f"Recommended: upload the {cur} statement for {period.day_before:%d/%m/%Y}. "
                                f"It may contain payments with {bank.COLUMNS['payment_date']} {period.start:%d/%m/%Y}.")

    # ---- reverse check (some banks): IremboPay payments of this bank that are on no statement
    if option(bank, "REVERSE_CHECK", False) and statements and pays:
        _reverse_check(period, bank, statements, pays, chk)

    # ---- the IremboPay payment file
    if not pays:
        chk.lines.append((False, "Payments: no IremboPay payment file added"))
        chk.blocking.append("Payments: no file added.")
    else:
        first = min(f.first for f in pays)
        last = max(f.last for f in pays)
        chk.lines.append((True, f"Payments: {first:%d/%m} to {last:%d/%m} ({sum(f.rows for f in pays):,} payments)"))
        if first > period.start or last < period.last_day:
            chk.warnings.append(f"The payment file looks short: it covers {first:%d/%m} to {last:%d/%m}, the period is "
                                f"{period.start:%d/%m} to {period.last_day:%d/%m}. Invoices outside it will show as not found.")
    return chk


def _reverse_check(period: Period, bank, statements: list, pays: list, chk: Check):
    """IremboPay payments made through this bank in the period that are on no uploaded statement."""
    p = pd.concat([f.data for f in pays], ignore_index=True)
    rows = bank.clean_statement(pd.concat([f.data for f in statements], ignore_index=True))
    if hasattr(bank, "find_invoices"):                    # rows without invoice found another way (BPR)
        rows = bank.find_invoices(rows, p)
    on_statement = set(rows["invoice_number"])
    p = p[p["payment_provider"].isin(bank_providers(bank)) & p["payment_status"].isin(bank.COMPLETED_STATUSES)
          & p["paid_at"].between(period.start, period.end, inclusive="left")]
    missing = p[~p["invoice_number"].isin(on_statement)].drop_duplicates("invoice_number").sort_values("paid_at")
    n = p["invoice_number"].nunique()
    title = f"Paid in IremboPay through {bank.NAME}"
    if missing.empty:
        chk.lines.append((True, f"{title}: all {n:,} are in the {bank.NAME} statements"))
        return
    # 26/10/2026: one line only (no warning, no list under it)
    k = len(missing)
    text = f"{title}: {k:,} of {n:,} {'is' if k == 1 else 'are'} not in the {bank.NAME} statements"
    chk.lines.append((False, text))


# =====================================================================================================================
# 7. REVIEW DECISIONS (step 4 of the app; also used by the developer notebook)
#    Rows not found in IremboPay are reviewed by the user. The decisions are saved in output/.review_store.csv,
#    one line per bank row: bank | period | key | the user's choices. Decisions of other periods are suggested again.
# =====================================================================================================================
def store_path(base_dir: Path) -> Path:
    return Path(base_dir) / "output" / STORE_NAME


def key_of(df, bank=None):
    """Key of a bank row in the review store: currency | payment date | bank reference | invoice."""
    bank = get_bank(bank)
    c = bank.COLUMNS
    if option(bank, "REVIEW_KEY_WITHOUT_REFERENCE", False):          # Ecobank: the reference is not used
        reference = pd.Series("", index=df.index)
    else:
        reference = df[c["reference"]]
        reference = reference.astype(object).where(reference.notna(), "")   # no reference: blank, the row still has
    return (df["Currency"].astype(str) + "|"                                # a key and gets its decision
            + pd.to_datetime(df[c["payment_date"]]).dt.strftime("%Y-%m-%d") + "|"
            + reference.astype(str).str.strip() + "|"
            + df["invoice_number"].fillna("").astype(str).str.strip())


def is_yes(value):
    return str(value).strip().lower() in ("yes", "y")


def load_store(base_dir: Path, bank=None) -> pd.DataFrame:
    """All review decisions (bank=None) or those of one bank. Rows saved before banks existed are I&M."""
    migrate_review_files(base_dir)
    path = store_path(base_dir)
    if not path.exists():
        return pd.DataFrame(columns=STORE_COLS)
    store = pd.read_csv(path, dtype=str).fillna("")
    for c in STORE_COLS:
        if c not in store.columns:
            store[c] = ""
    store.loc[store["bank"].eq(""), "bank"] = "IM"
    # banks whose review ignores the reference: decisions saved with one still match (currency | date | | invoice)
    for code in [k for k, b in BANKS.items() if option(b, "REVIEW_KEY_WITHOUT_REFERENCE", False)]:
        rows = store["bank"].eq(code)
        if rows.any():                                    # currency | date (the first two parts) + the invoice from
            parts = store.loc[rows, "key"].str.split("|", n=2, expand=True)   # its own column: a reference or an
            if parts.shape[1] == 3:                                            # invoice may itself contain "|"
                store.loc[rows, "key"] = (parts[0] + "|" + parts[1] + "||"
                                          + store.loc[rows, "invoice_number"].astype(str).str.strip())
    store = store[STORE_COLS]
    return store if bank is None else store[store["bank"].eq(get_bank(bank).CODE)]


def _write_store(base_dir: Path, store: pd.DataFrame, what="save again"):
    """Save the whole review store (a clear message if the file is open in Excel)."""
    path = store_path(base_dir)
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        store.to_csv(path, index=False)
    except PermissionError:
        raise ReconError(f"The review store is open in another program. Close it and {what}.") from None


def save_decisions(base_dir: Path, period: Period, review: pd.DataFrame, bank=None):
    """Write the decisions of one period (replaces that period's earlier decisions for this bank)."""
    bank = get_bank(bank)
    store = load_store(base_dir)
    rows = review.copy()
    rows["key"] = key_of(rows, bank)
    rows["bank"] = bank.CODE
    rows["period"] = period.store_key
    rows["saved_at"] = pd.Timestamp.now().strftime("%Y-%m-%d %H:%M:%S")
    rows["invoice_number"] = rows["invoice_number"].fillna("").astype(str)
    for c in EDIT_COLS:
        rows[c] = rows[c].fillna("").astype(str).replace({"None": "", "nan": ""})
    rows = rows[rows[EDIT_COLS].apply(lambda r: any(v.strip() for v in r), axis=1)]     # skip rows left empty
    store = store[~(store["bank"].eq(bank.CODE) & store["period"].eq(period.store_key))]
    _write_store(base_dir, pd.concat([store, rows[STORE_COLS]], ignore_index=True))


def migrate_review_files(base_dir: Path):
    """One time: copy the decisions of the old Excel review files into the store, then move those files away."""
    base_dir = Path(base_dir)
    path = store_path(base_dir)
    old_files = sorted((base_dir / "output").glob("*/review/IM_missing_review_*.xlsx"))
    if path.exists() or not old_files:
        return
    rows = []
    for f in old_files:
        try:
            h = pd.read_excel(f, dtype=str).fillna("")
        except Exception:
            continue
        if h.empty:
            continue
        for c in EDIT_COLS:
            if c not in h.columns:
                h[c] = ""
        h = h[h[EDIT_COLS].apply(lambda r: any(str(v).strip() for v in r), axis=1)].copy()
        if h.empty:
            continue
        h["key"] = key_of(h, "IM")
        h["bank"] = "IM"
        h["period"] = f"{f.parent.parent.name}:{f.stem.replace('IM_missing_review_', '')}"
        h["saved_at"] = pd.Timestamp(f.stat().st_mtime, unit="s").strftime("%Y-%m-%d %H:%M:%S")
        rows.append(h[STORE_COLS])
    path.parent.mkdir(parents=True, exist_ok=True)
    (pd.concat(rows, ignore_index=True) if rows else pd.DataFrame(columns=STORE_COLS)).to_csv(path, index=False)
    keep = base_dir / "output" / ".old_review_files"
    for f in old_files:
        target = keep / f.parent.parent.name
        target.mkdir(parents=True, exist_ok=True)
        f.rename(target / f.name)
    for d in {f.parent for f in old_files}:
        try:
            d.rmdir()
        except OSError:
            pass


def _period_title(store_key: str) -> str:
    """'month:01-08-2026_to_01-09-2026' -> 'August 2026 review' (shown in the 'Suggested from' column)."""
    try:
        kind, label = store_key.split(":")
        start = pd.to_datetime(label.split("_to_")[0], format="%d-%m-%Y")
        return f"{start:%B %Y} review" if kind == "month" else f"{kind} review of {start:%d/%m/%Y}"
    except Exception:
        return store_key


def review_rows(base_dir: Path, period: Period, not_found: pd.DataFrame, bank=None):
    """Rows to review for this period: this period's decisions, or suggestions from other periods (same bank).
    Returns (review table, number of suggestions)."""
    bank = get_bank(bank)
    store = load_store(base_dir, bank)
    review = not_found[bank_cols(bank)].copy().reset_index(drop=True)
    review["_key"] = key_of(review, bank)
    for c in EDIT_COLS + [SUGGEST_COL]:
        review[c] = ""
    mine = store[store["period"] == period.store_key].drop_duplicates("key", keep="last").set_index("key")
    others = store[(store["period"] != period.store_key) & store["Reviewed (Yes/No)"].map(is_yes).astype(bool)]
    by_key = others.drop_duplicates("key", keep="last").set_index("key")
    by_inv = (others[others["invoice_number"].ne("")]
              .drop_duplicates("invoice_number", keep="last").set_index("invoice_number"))
    suggested = 0
    no_currency = option(bank, "CURRENCY_FROM_IREMBOPAY", False)
    for i, r in review.iterrows():
        key = r["_key"]
        if key not in mine.index and no_currency:         # saved while its currency was still unknown (blank)
            blank = "|" + key.split("|", 1)[1]
            key = blank if blank in mine.index else key
        if key in mine.index:                             # 1. decided for this period: use it as it is
            for c in EDIT_COLS:
                review.at[i, c] = mine.loc[key, c]
            continue
        inv = str(r["invoice_number"]) if pd.notna(r["invoice_number"]) else ""
        source = by_key.loc[r["_key"]] if r["_key"] in by_key.index else (by_inv.loc[inv] if inv in by_inv.index else None)
        if source is not None:                            # 2. decided in another period: a suggestion to confirm
            for c in EDIT_COLS:
                if c != "Reviewed (Yes/No)":
                    review.at[i, c] = source[c]
            review.at[i, SUGGEST_COL] = _period_title(source["period"])
            suggested += 1
    return review, suggested


def apply_review(merged: pd.DataFrame, not_found: pd.DataFrame, payments: pd.DataFrame, review: pd.DataFrame,
                 bank=None):
    """Apply the review decisions to the matched rows (same logic in the app and the notebook).
    Returns (merged, not_found_with_status, counts, messages)."""
    bank = get_bank(bank)
    amount = bank.COLUMNS["amount"]
    merged = merged.copy()
    current = merged["period"].eq("current")
    merged["originally_not_found"] = ~merged["found_in_irembopay"]
    to_review = merged["_to_review"] if "_to_review" in merged else merged["originally_not_found"]
    choose_currency = option(bank, "CURRENCY_FROM_IREMBOPAY", False)
    for c in ("review_status", "review_table", "review_comment", "sla_override"):
        merged[c] = ""
    merged["_key"] = key_of(merged, bank)
    pay_by_invoice = payments.set_index("invoice_number")
    payment_fields = [c for c in payments.columns if c != "invoice_number"]
    counts, messages, exclude_keys = {"Not reviewed": 0}, [], []

    for _, r in review.iterrows():
        rows = current & to_review & merged["_key"].eq(r["_key"])
        if not is_yes(r["Reviewed (Yes/No)"]):
            merged.loc[rows, "review_status"] = "Not reviewed"
            counts["Not reviewed"] += 1
            continue
        table_choice = str(r["Report table"]).strip()
        if table_choice and table_choice not in TABLE_CHOICES:
            table_choice = ""
        correct_inv = str(r["Correct invoice number"]).strip()
        if choose_currency and table_choice != "Exclude":   # no currency on the statement: IremboPay's, or chosen
            known = merged.loc[rows, "Currency"].fillna("").astype(str).ne("").all()
            chosen = str(r.get(CURRENCY_CHOSEN, "")).strip().upper()
            if correct_inv and correct_inv in pay_by_invoice.index:
                merged.loc[rows, "Currency"] = pay_by_invoice.loc[correct_inv, "currency"]
            elif chosen in bank.CURRENCIES:
                merged.loc[rows, "Currency"] = chosen
            elif not known:                                 # reviewed, but its currency is still unknown
                merged.loc[rows, "review_status"] = "Not reviewed"
                counts["Not reviewed"] += 1
                counts["Currency to choose"] = counts.get("Currency to choose", 0) + 1
                continue
        manual = {"payment_status": str(r["payment_status"]).strip().upper(),
                  "payment_provider": str(r["payment_provider"]).strip(),
                  "paid_at": str(r["paid_at (YYYY-MM-DD HH:MM:SS)"]).strip(),
                  "name": str(r["Merchant name"]).strip()}

        # ---- a corrected invoice number: take its IremboPay record
        if correct_inv:
            if correct_inv in pay_by_invoice.index:
                record = pay_by_invoice.loc[correct_inv]
                for col in payment_fields:
                    merged.loc[rows, col] = record[col]
                merged.loc[rows, "found_in_irembopay"] = True
            else:
                messages.append(("warning", f"Invoice {correct_inv} (typed in the review) is not in the payment file; "
                                            "the typed details are used instead."))
        # ---- details typed by the user (status, provider, paid_at, merchant)
        for col, value in manual.items():
            if value:
                merged.loc[rows, col] = pd.to_datetime(value, errors="coerce") if col == "paid_at" else value
        if any(manual.values()):
            merged.loc[rows, "found_in_irembopay"] = True
            merged.loc[rows, "amount"] = merged.loc[rows, "amount"].fillna(merged.loc[rows, amount])
            merged.loc[rows, "currency"] = merged.loc[rows, "currency"].fillna(merged.loc[rows, "Currency"])

        # ---- the Report table chosen: puts the row in the right summary table
        if table_choice == "Exclude":
            exclude_keys.append(r["_key"])
        elif table_choice == "Notification received":
            merged.loc[rows, "payment_provider"] = bank.PROVIDER
            if manual["payment_status"] not in bank.COMPLETED_STATUSES:
                merged.loc[rows, "payment_status"] = "PAID"
            merged.loc[rows, "found_in_irembopay"] = True
        elif table_choice == "Pending notification":
            merged.loc[rows, "payment_provider"] = bank.PROVIDER
            # always NEW (20/10/2026): "expired" is not a payment status - an expired invoice was never paid, so its
            # payment status is NEW; the expiry is shown as a note (review status / comment)
            merged.loc[rows, "payment_status"] = "NEW"
            merged.loc[rows, "found_in_irembopay"] = True
        elif table_choice == "Paid with other channel":
            # 27/10/2026: the provider IremboPay recorded is kept (its origin); the one typed in the review, else
            # "OTHER", only when IremboPay has none
            provider = manual["payment_provider"]
            recorded = merged.loc[rows, "payment_provider"]
            typed = provider if provider and provider not in bank_providers(bank) else "OTHER"
            merged.loc[rows, "payment_provider"] = recorded.where(recorded.notna() & recorded.astype(str).str.strip()
                                                                  .ne(""), typed)
            if manual["payment_status"] in ("", "NEW"):
                merged.loc[rows, "payment_status"] = "PAID"
            merged.loc[rows, "found_in_irembopay"] = True
        elif table_choice == "Failed notification":
            for col in payment_fields:
                merged.loc[rows, col] = None
            merged.loc[rows, "found_in_irembopay"] = False
        if table_choice in ("Notification received", "Pending notification", "Paid with other channel"):
            merged.loc[rows, "amount"] = merged.loc[rows, "amount"].fillna(merged.loc[rows, amount])
            merged.loc[rows, "currency"] = merged.loc[rows, "currency"].fillna(merged.loc[rows, "Currency"])

        # ---- what the row became (shown in the report), and an SLA decided by hand
        expired = (manual["payment_status"] == "EXPIRED"
                   or "expired" in str(r["Comment"]).lower())    # IremboPay checkout: the invoice expired unpaid
        if table_choice == "Pending notification" and expired:
            status = "Pending notification (invoice expired)"
        elif table_choice == "Exclude":
            status = EXCLUDED
        elif table_choice:
            status = table_choice
        elif merged.loc[rows, "found_in_irembopay"].any():
            status = f"Matched to invoice {correct_inv}" if correct_inv in pay_by_invoice.index else "Resolved manually"
        else:
            status = "Failed notification"
        counts[status] = counts.get(status, 0) + 1
        merged.loc[rows, "review_status"] = status
        merged.loc[rows, "review_table"] = table_choice
        merged.loc[rows, "review_comment"] = r["Comment"]
        sla = str(r["SLA violated (Yes/No)"]).strip().lower()
        if sla in ("yes", "y", "no", "n"):
            merged.loc[rows, "sla_override"] = "SLA violated" if sla.startswith("y") else "On time"
    merged["paid_at"] = pd.to_datetime(merged["paid_at"], errors="coerce")
    if counts.get("Currency to choose"):
        messages.append(("warning", f"{counts['Currency to choose']} credit(s) marked reviewed have no currency (not in "
                                    "IremboPay): choose it in the Currency column, then save again."))

    nf = not_found.copy()
    nf["_key"] = key_of(nf, bank)
    info = merged.drop_duplicates("_key").set_index("_key")[["review_status", "review_comment"]]
    nf = nf.join(info, on="_key").drop(columns="_key")
    if exclude_keys:                  # "Exclude" (26/10/2026): the row stays in the statement, in no table
        messages.append(("info", f"{len(exclude_keys)} row(s) excluded, as chosen in the review: they stay in the "
                                 "statement with their comment, but are counted in no table."))
    return merged, nf, counts, messages


# =====================================================================================================================
# 8. COMPARISON WITH THE PREVIOUS REPORT ("Compare with previous period" in step 1)
#    The previous report of the same bank is read back; each table gets a "VS prev <unit> (%)" column.
# =====================================================================================================================
def read_report_counts(path: Path, sheet: str, bank=None):
    """({table title: {"RWF": n, "USD": n, "Total": n}}, heading, rows not reviewed) from the first block of a
    summary sheet of a report made by this app or the reconciliation notebooks (whichever column their tables are
    in: the title text is found first, then its numbers are read relative to it). The notebooks never write
    "Pending notifications" to their Summary itself - derived here from their own "Pending notification" sheet
    (counted; 0 when it is not there), the same way engine.read_notebook_report does. "Overall notifications
    received by Irembo" is always derived too (Processed - Pending - Duplicate, CLAUDE.md's own definition), never
    read from the file: some of the notebooks' own reports do carry a table under that exact title, but it is a
    different count (every payment Irembo's checkout logged, not the ones it shows as paid) - approved 2026-10-27,
    found while wiring up "Compare with previous period" against the notebooks' reports."""
    bank = get_bank(bank)
    wb = load_workbook(path, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise ReconError(f"{path.name}: this report has no {sheet} sheet.")
        ws = wb[sheet]
        heading = ws.cell(1, 1).value or ""
        counts = {}
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30)):
            for c in row:
                if c.value in TABLE_TITLES and c.value != "Overall notifications received by Irembo" and c.value not in counts:
                    vals = {}
                    for i, cur in enumerate(bank.CURRENCIES + ["Total"]):
                        v = ws.cell(c.row + 2 + i, c.column + 1).value
                        vals[cur] = int(v) if isinstance(v, (int, float)) else 0
                    counts[c.value] = vals
        if "Pending notifications" not in counts:
            pending = _notebook_sheet_counts(wb, "Pending notification", bank.CURRENCIES)
            counts["Pending notifications"] = ({"Total" if pd.isna(c) else c: int(n) for c, n in
                                               zip(pending["Currency"], pending["Number of Transactions"])}
                                               if pending is not None
                                               else {**{cur: 0 for cur in bank.CURRENCIES}, "Total": 0})
        if "Overall processed payments" in counts:
            counts["Overall notifications received by Irembo"] = {
                cur: counts["Overall processed payments"].get(cur, 0) - counts["Pending notifications"].get(cur, 0)
                     - counts.get("Duplicate Payment", {}).get(cur, 0)
                for cur in bank.CURRENCIES + ["Total"]}
        not_reviewed = max(wb[NOT_FOUND_SHEET].max_row - 1, 0) if NOT_FOUND_SHEET in wb.sheetnames else 0
    finally:
        wb.close()
    missing = [t for t in TABLE_TITLES if t not in counts]
    if missing:
        raise ReconError(f"{path.name}: this report does not have the expected tables ({', '.join(missing)}).")
    return counts, heading, not_reviewed


def find_previous_report(base_dir: Path, period: Period, bank=None):
    """(path or None, previous Period): the previous report of the same bank."""
    rename_old_reports(base_dir)
    prev = period.previous()
    path = report_path(base_dir, prev, bank)
    return (path if path.exists() else None), prev


def save_uploaded_report(base_dir: Path, period: Period, name: str, content: bytes, bank=None,
                         root: Path | None = None) -> Path:
    """Check an uploaded report for `period` and keep it in the report store (root: see report_path)."""
    tmp = (Path(root) if root is not None else Path(base_dir) / "output") / f".upload_{file_digest(content)}.xlsx"
    tmp.parent.mkdir(parents=True, exist_ok=True)
    tmp.write_bytes(content)
    try:
        wb = load_workbook(tmp, read_only=True)
        sheets = set(wb.sheetnames)
        wb.close()
        mine = get_bank(bank)
        made_here = [kind for kind, sheet in SUMMARY_SHEET.items() if sheet in sheets]
        if made_here and is_no_transaction_report(tmp, period):  # a No-Transaction Report of this bank and period
            wb = load_workbook(tmp, read_only=True, data_only=True)
            ws = wb[SUMMARY_SHEET[period.report_type]]
            title, dates = str(ws.cell(1, 1).value or ""), str(ws.cell(2, 1).value or "")
            wb.close()
            if not title.startswith(mine.FULL_NAME):
                raise ReconError(f"{name}: this No-Transaction Report is not for {mine.NAME} ({title}).")
            if f"{period.start:%d/%m/%Y}" not in dates:
                raise ReconError(f"{name}: this No-Transaction Report is for '{dates}', not for {period.name}.")
            target = report_path(base_dir, period, bank, root)
            target.parent.mkdir(parents=True, exist_ok=True)
            tmp.replace(target)
            return target
        if not made_here or not any(b.STATEMENT_SHEET in sheets for b in BANKS.values()):
            raise ReconError(f"{name}: this is not a reconciliation report made by this app.")
        if mine.STATEMENT_SHEET not in sheets:                       # 1. the right bank
            other = next(b.NAME for b in BANKS.values() if b.STATEMENT_SHEET in sheets)
            raise ReconError(f"{name}: this report is for {other}, not {mine.NAME}.")
        if SUMMARY_SHEET[period.report_type] not in sheets:          # 2. the right kind (day / week / month)
            kinds = {"day": "a day", "week": "a week", "month": "a month"}
            is_kind = "month" if "month" in made_here else made_here[-1]
            raise ReconError(f"{name}: this is {kinds[is_kind]} report; {period.name} needs {kinds[period.report_type]} "
                             "report.")
        _, heading, _ = read_report_counts(tmp, SUMMARY_SHEET[period.report_type], bank)
        if f"{period.start:%d/%m/%Y}" not in heading:                # 3. the right period
            raise ReconError(f"{name}: this report is for '{heading}', not for {period.name}.")
    except ReconError as err:
        tmp.unlink(missing_ok=True)
        raise ReconError(str(err).replace(tmp.name, name)) from None
    except Exception:
        tmp.unlink(missing_ok=True)
        raise ReconError(f"{name}: this is not a reconciliation report made by this app.") from None
    target = report_path(base_dir, period, bank, root)
    target.parent.mkdir(parents=True, exist_ok=True)
    tmp.replace(target)
    return target


def counts_of(tables: dict) -> dict:
    """{table title: {currency or "Total": number of transactions}} of one block."""
    return {title: {("Total" if pd.isna(c) else c): int(n) for c, n in zip(t["Currency"], t["Number of Transactions"])}
            for title, t in tables.items()}


def merchant_counts(table) -> dict:
    """{(merchant, currency): number of transactions} of a Payment per merchant table ("All" included)."""
    if table is None:
        return {}
    return {(str(m), str(c)): int(n) for m, c, n in zip(table["Merchant name"], table["Currency"],
                                                       table["Number of Transactions"])}


def add_vs_merchants(table: pd.DataFrame, prev, unit: str, prev_no_tx: bool = False) -> pd.DataFrame:
    """Add 'VS prev <unit> (%)' to a Payment per merchant table: the change of each merchant (same currency).
    prev None = nothing to compare (empty column); a merchant missing before (or a previous report without
    the table, prev {}) = "N/A"."""
    t = table.copy()
    col = f"VS prev {unit} (%)"
    if prev is None:
        t[col] = None
        return t
    out = []
    for m, c, n in zip(t["Merchant name"], t["Currency"], t["Number of Transactions"]):
        before = prev.get((str(m), str(c)), 0)
        out.append(_from_zero(n) if prev_no_tx and not before
                   else "N/A" if not before else f"{(n - before) / before * 100:.2f}%")
    t[col] = out
    return t


def _from_zero(n) -> str:
    """The change from a period with no transactions: 'No change' (still none) or '+100.00%' (some now)."""
    return "No change" if not n else "+100.00%"


def add_vs_column(tables: dict, prev_counts, unit: str, prev_no_tx: bool = False) -> dict:
    """Add 'VS prev <unit> (%)' (change in the number of transactions); empty when there is nothing to compare.
    prev_no_tx: the previous period was a No-Transaction Report (all 0): 'No change' or '+100.00%', never N/A."""
    col = f"VS prev {unit} (%)"
    out = {}
    for title, t in tables.items():
        t = t.copy()
        if prev_counts is None:
            t[col] = None
        else:
            prev = prev_counts[title]
            t[col] = [(_from_zero(n) if prev_no_tx and not prev["Total" if pd.isna(c) else c]
                       else "N/A" if not prev["Total" if pd.isna(c) else c]
                       else f"{(n - prev['Total' if pd.isna(c) else c]) / prev['Total' if pd.isna(c) else c] * 100:.2f}%")
                      for c, n in zip(t["Currency"], t["Number of Transactions"])]
        out[title] = t
    return out


# =====================================================================================================================
# 9. THE RECONCILIATION RUN ("Run reconciliation" in step 3)
#    run() goes through steps A to I below. It stops after step D when rows still need a review.
# =====================================================================================================================
@dataclass
class RunResult:
    period: Period
    bank: str = ""
    messages: list = field(default_factory=list)      # (level, text): "info", "warning" or "success"
    bank_rows: int = 0                                # statement rows in the period
    not_found: int = 0                                # rows not found in IremboPay
    to_review: int = 0                                # rows in Exception Resolution: not found + found but to check
    currency_to_choose: int = 0                       # reviewed rows whose currency is still to choose (BPR)
    review: pd.DataFrame | None = None                # the rows to review (step 4 of the app)
    not_reviewed: int = 0
    stopped_for_review: bool = False                  # True: no report yet, some rows are not reviewed
    blocks: list = field(default_factory=list)        # [(sheet, heading, tables)]
    report_file: Path | None = None
    timings: dict = field(default_factory=dict)       # seconds per step (written in the run log)
    compared_with: str = ""
    ran_at: pd.Timestamp = field(default_factory=pd.Timestamp.now)       # when the run was made
    by_currency: dict = field(default_factory=dict)   # {currency: {"rows", "matched", "not_found", "amount"}}
    review_counts: dict = field(default_factory=dict) # what the reviewed rows became: {"Pending notification": 6, ...}
    sheets: dict = field(default_factory=dict)        # the detail sheets of the report (used by the custom report)
    daily: dict = field(default_factory=dict)         # {day: counts_of(the 5 tables of that day)} (custom charts)
    merchants: dict = field(default_factory=dict)     # {block heading: Payment per merchant table}
    previous_counts: dict | None = None               # "Compare": counts_of() the previous report (Slack message)

    @property
    def reconciled_on(self) -> pd.Timestamp:
        """The day of the run."""
        return self.ran_at.normalize()

    @property
    def matched(self) -> int:
        """Bank rows found in the IremboPay file by invoice number (before the review)."""
        return self.bank_rows - self.not_found

    def add(self, level, text):
        self.messages.append((level, text))


def run(base_dir: Path, period: Period, files: list, sla_days: int = 0, compare: bool = False,
        require_review: bool = True, bank=None, report_dir: Path | None = None,
        previous_report: Path | None = None, checkout: dict | None = None) -> RunResult:
    """Reconcile the uploaded files for one bank and period, and write the report.
    report_dir: where the Excel is written (default output/; the app gives a temporary folder).
    previous_report: the report compared with (default: looked for in output/).
    checkout: {invoice: IremboPay checkout answer} of every invoice, looked up before the run (banks with
    CHECKOUT_ALL: Ecobank, BPR); kept in the report's Payment details ("IremboPay check").
    Steps: A payments  B statement  C match  D review  E SLA  F summary tables  G totals check  H comparison
    I report + run log."""
    bank = get_bank(bank)
    c = bank.COLUMNS
    posted, paid_on, desc, ref, amount = (c["posted_date"], c["payment_date"], c["description"],
                                          c["reference"], c["amount"])
    base_dir = Path(base_dir)
    output_dir = base_dir / "output"
    start, end = period.start, period.end
    res = RunResult(period, bank.CODE)
    t0 = time.perf_counter()

    not_ended = period_not_ended(period, res.reconciled_on)
    if not_ended:
        raise ReconError(not_ended.replace("\n", " "))
    statements = [f.data for f in files if f.kind == "bank"]
    pays = [f.data for f in files if f.kind == "payments"]
    if not statements:
        raise ReconError("No bank file was added.")
    if not pays:
        raise ReconError("No IremboPay payment file was added.")

    # ---- A. payments: one per invoice (the successful transaction, else the latest one) ----
    payments = pd.concat(pays, ignore_index=True)
    if option(bank, "ALL_PROVIDER_REFERENCES", False):     # every IremboPay transaction of an invoice (BPR: which
        refs = (payments.assign(_r=payments["provider_reference"].fillna("").astype(str).str.strip())   # credits it
                .groupby("invoice_number")["_r"].agg(lambda r: "|".join(sorted(set(r) - {""}))))       # recorded)
        payments["_provider_refs"] = payments["invoice_number"].map(refs).fillna("")
    payments["_successful"] = payments["transaction_status"].astype(str).str.upper().eq("SUCCESSFUL")
    payments = (payments.sort_values(["invoice_number", "_successful", "created_at"], ascending=[True, False, False])
                .drop_duplicates("invoice_number", keep="first").drop(columns="_successful").reset_index(drop=True))

    # ---- B. bank statement: all files together, repeated rows dropped (by default the period is set by the
    #         payment date); some banks keep one row per invoice ----
    statement = bank.clean_statement(pd.concat(statements, ignore_index=True))
    statement = statement.drop_duplicates(
        subset=option(bank, "DUPLICATE_KEY", None) or ["Currency"] + present(ref, posted, paid_on, desc, amount)
    ).reset_index(drop=True)
    if hasattr(bank, "find_invoices"):                    # some banks: rows without invoice found another way (BPR)
        statement = bank.find_invoices(statement, payments)
    attempts = statement                                  # every statement row (some banks keep one per invoice)
    if hasattr(bank, "one_row_per_invoice"):
        statement = bank.one_row_per_invoice(statement, payments).reset_index(drop=True)
    res.timings["read"] = time.perf_counter() - t0

    # ---- C. match: every statement row with its IremboPay payment (by invoice number) ----
    t = time.perf_counter()
    # a payment file sometimes carries the bank's own columns too (e.g. someone re-uploads a report's own Payment
    # details sheet by mistake, which already has Request Date next to invoice_number): kept as-is, such a column
    # would collide with the statement's own in the merge below and silently become "<name>_x" / "<name>_y",
    # breaking a bank's own column-by-name rules (KeyError, e.g. "Request Date" - found 2026-10-28). The statement
    # always wins; this never drops a column the payments side is actually read by name for (present(...)).
    clash = (set(payments.columns) & set(statement.columns)) - {"invoice_number"}
    if clash:
        payments = payments.drop(columns=list(clash))
    merged = statement.merge(payments, on="invoice_number", how="left", indicator=True)
    merged["found_in_irembopay"] = merged["_merge"].eq("both")
    merged = merged.drop(columns="_merge")
    if option(bank, "CURRENCY_FROM_IREMBOPAY", False):     # no currency on the statement: the IremboPay one
        merged["Currency"] = merged["currency"].astype(object).where(merged["found_in_irembopay"]
                                                                     & merged["currency"].notna(), "")
        if hasattr(bank, "currency_not_found"):            # a credit IremboPay's file does not have: the bank's rule
            lost = ~merged["found_in_irembopay"]                                  # (the reviewer can change it)
            merged.loc[lost, "Currency"] = bank.currency_not_found(merged[lost])
    if hasattr(bank, "in_period"):                        # the bank's own rule
        in_period = bank.in_period(merged, attempts, start, end)
    else:                                                 # default: payment date inside the period
        in_period = merged[paid_on].between(start, end, inclusive="left")
    merged["period"] = None
    merged.loc[in_period, "period"] = "current"
    current = merged["period"].eq("current")
    if not current.any():
        raise NoTransactionsError(f"The bank files have no Irembo payments with a {paid_on} in {period.name}.")
    no_invoice = merged.loc[current, "invoice_number"].isna().sum()
    if no_invoice:
        res.add("warning", f"{no_invoice} bank row(s) have no invoice number in the description.")
    res.bank_rows = int(current.sum())
    flagged = pd.Series(False, index=merged.index)
    reasons = pd.Series("", index=merged.index, dtype=object)
    if hasattr(bank, "review_reasons"):                   # found, but to be checked (e.g. BPR: another provider)
        reasons[current] = bank.review_reasons(merged[current])
        flagged = current & merged["found_in_irembopay"] & reasons.ne("")
    merged["_to_review"] = (current & ~merged["found_in_irembopay"]) | flagged
    # the extra credits of an invoice credited more than once (BPR): never the one counted in the summary tables
    merged["_credited_again"] = reasons.astype(str).str.startswith("Invoice credited again")
    no_invoice_pending = current & merged["invoice_number"].isna() & option(bank, "NO_INVOICE_PENDING", False)
    for col in option(bank, "ALWAYS_REVIEW_COLUMNS", []):   # e.g. a credit cancelled by a debit: still reviewed
        if col in merged:
            no_invoice_pending &= merged[col].fillna("").astype(str).eq("")
    merged.loc[no_invoice_pending, "_to_review"] = False   # BPR: no invoice number -> Pending, never reviewed
    per_cur = merged[current].groupby("Currency").agg(rows=("found_in_irembopay", "size"),
                                                      matched=("found_in_irembopay", "sum"), amount=(amount, "sum"))
    order = [cur for cur in bank.CURRENCIES if cur in per_cur.index] + [cur for cur in per_cur.index
                                                                         if cur not in bank.CURRENCIES]
    res.by_currency = {(cur or "Not in IremboPay"): {"rows": int(per_cur.at[cur, "rows"]), "matched": int(per_cur.at[cur, "matched"]),
                             "not_found": int(per_cur.at[cur, "rows"] - per_cur.at[cur, "matched"]),
                             "amount": float(per_cur.at[cur, "amount"])} for cur in order}
    not_found = merged[merged["_to_review"]].copy()
    not_found["Bank channel"] = bank.channel(not_found)
    if hasattr(bank, "not_found_reason"):
        not_found["Reason"] = bank.not_found_reason(not_found)
    else:
        not_found["Reason"] = ["No invoice number in description" if pd.isna(inv)
                               else "Invoice not in IremboPay payment file" for inv in not_found["invoice_number"]]
    found_rows = not_found["found_in_irembopay"].astype(bool)
    not_found.loc[found_rows, "Reason"] = reasons[found_rows[found_rows].index]
    again = ~found_rows & reasons.reindex(not_found.index).fillna("").astype(str).str.startswith("Invoice credited again")
    if again.any():
        not_found.loc[again, "Reason"] = [f"{a} ({b})" for a, b in zip(reasons.reindex(not_found.index)[again],
                                                                       not_found.loc[again, "Reason"])]
    not_found = not_found[bank_cols(bank)].sort_values(paid_on)
    res.not_found = int((current & ~merged["found_in_irembopay"]).sum())
    res.to_review = len(not_found)
    res.timings["match"] = time.perf_counter() - t

    # ---- D. review: apply the saved decisions; stop here when some rows are not reviewed yet ----
    t = time.perf_counter()
    review, suggested = review_rows(base_dir, period, not_found, bank)
    if suggested:
        res.add("info", f"{suggested} invoice(s) were already decided in an earlier review: the decision is filled in "
                        "as a suggestion. Check it and set Reviewed = Yes.")
    merged, not_found, counts, msgs = apply_review(merged, not_found, payments, review, bank)
    if no_invoice_pending.any():                          # BPR: the credits with no invoice number are pending
        auto = merged["period"].eq("current") & merged["invoice_number"].isna() & ~merged["_to_review"]
        merged.loc[auto, "payment_status"] = "NEW"
        merged.loc[auto, "review_status"] = "Pending notification (no invoice number)"
        res.add("info", f"{int(auto.sum())} credit(s) with no invoice number are in Pending notification.")
    res.messages += msgs
    res.review = review.drop(columns="_key").reset_index(drop=True)
    res.not_reviewed = counts["Not reviewed"]
    res.currency_to_choose = counts.pop("Currency to choose", 0)          # BPR: reviewed rows without a currency
    res.review_counts = {k: v for k, v in counts.items() if v and k != "Not reviewed"}
    res.timings["review"] = time.perf_counter() - t
    if require_review and res.not_reviewed:
        res.stopped_for_review = True
        _log_run(output_dir, res, sla_days)
        return res

    # ---- E. SLA (the bank's rule) ----
    t = time.perf_counter()
    merged["paid_date"] = merged["paid_at"].dt.normalize()
    sla_col = option(bank, "SLA_DIFF_COLUMN", "sla_days_difference")
    merged[sla_col], merged["sla_status"] = bank.sla(merged, sla_days)

    # ---- F. summary: flags per row (the bank's calculations) + a reason when a row is not a clean match ----
    rep = (merged[merged["period"] == "current"].sort_values(present(paid_on, posted, ref))
           .reset_index(drop=True))
    has_invoice = rep["invoice_number"].notna()
    if checkout is not None:                               # what IremboPay's checkout said, before the run
        rep["IremboPay check"] = rep["invoice_number"].map(lambda i: checkout.get(str(i), "") if pd.notna(i) else "")
    rep["times_on_statement"] = rep.groupby("invoice_number")["invoice_number"].transform("size")
    # the invoice is the unique identifier: an invoice credited more than once counts once (the credit IremboPay
    # recorded when the bank says which one, otherwise the first); the others are extra credits
    order = rep.sort_values(["_credited_again"], kind="stable") if "_credited_again" in rep else rep
    extra = order.index[has_invoice[order.index] & order.duplicated("invoice_number", keep="first")]
    rep["is_extra_credit"] = rep.index.isin(extra)
    if "review_table" in rep:          # 26/10/2026: an extra credit in no table stays in the statement, with a comment
        loose = rep["is_extra_credit"] & ~rep["review_table"].fillna("").isin(list(EXTRA_DECIDED) + ["Exclude"])
        status = rep["review_status"].fillna("").astype(str)
        rep.loc[loose, "review_status"] = status[loose].map(lambda v: f"{v} - {NOT_COUNTED.lower()}" if v
                                                            else NOT_COUNTED)
    rep["paid_other_channel"] = bank.duplicate(rep)
    rep["notification_received"] = bank.received(rep)
    rep["is_pending"] = bank.pending(rep)
    found_ = rep["found_in_irembopay"]
    provider_ = rep["payment_provider"].astype(object).where(rep["payment_provider"].notna(), "none").astype(str)
    status_ = rep["payment_status"].astype(object).where(rep["payment_status"].notna(), "none").astype(str)
    amount_diff = rep["amount"].notna() & (rep["amount"] - rep[amount]).abs().gt(0.005)
    currency_diff = rep["currency"].notna() & rep["currency"].ne(rep["Currency"])
    parts = [
        ("Other channel (" + provider_ + ")").where(found_ & ~rep["payment_provider"].isin(bank_providers(bank)), ""),
        ("Not settled (" + status_ + ")").where(found_ & ~rep["payment_status"].isin(bank.COMPLETED_STATUSES), ""),
        ("Amount difference (bank " + rep[amount].map("{:,.2f}".format) + " vs IremboPay "
         + rep["amount"].fillna(0).map("{:,.2f}".format) + ")").where(found_ & amount_diff, ""),
        ("Currency difference (bank " + rep["Currency"].astype(str) + " vs IremboPay "
         + rep["currency"].astype(object).where(rep["currency"].notna(), "").astype(str) + ")").where(found_ & currency_diff, ""),
    ]
    combined = parts[0]
    for extra in parts[1:]:
        combined = combined + "; " + extra
    combined = combined.str.replace("(?:; )+", "; ", regex=True).str.strip("; ")
    rep["Reason"] = combined.where(found_, "Not found in IremboPay")

    def table_rows(a, b):
        """Rows of each of the 5 tables for [a, b): the bank's own rule, or by payment date (default).
        Rule (26/10/2026): Processed = every credit on the statement whose bank date is in [a, b), each invoice once
        (a second credit goes to Exception Resolution). Notification received = processed, paid (paid_at) through
        this bank; Duplicate Payment = paid through another provider; Pending = in IremboPay, not paid. What the
        rules do not cover is decided in Exception Resolution. SLA violated = the late ones of Processed."""
        window = (rep[paid_on] >= a) & (rep[paid_on] < b)        # the bank date (Value Date, VALUE_DATE, Request Date)
        if hasattr(bank, "tables"):
            return with_extra_decisions(rep, bank.tables(rep, attempts, a, b), window)
        processed = window & ~rep["is_extra_credit"]              # every credit on the statement, each invoice once
        received = processed & rep["notification_received"]
        duplicate = processed & rep["paid_other_channel"] & ~received
        pending = processed & rep["is_pending"] & ~received & ~duplicate
        return with_extra_decisions(rep, {"processed": processed, "received": received,
                                          "sla": processed & rep["sla_status"].eq("SLA violated"), "pending": pending,
                                          "duplicate": duplicate}, window)

    def build_summary5(a, b):
        """The 5 summary tables for [a, b): one row per currency + a total row (count only, no amount).
        % = share of the processed payments of the same currency."""
        m = table_rows(a, b)
        df = rep[m["processed"]]
        processed = df.groupby("Currency").size()

        def table(part, with_pct=True):
            rows = []
            for cur in bank.CURRENCIES:
                sub = part[part["Currency"] == cur]
                base = processed.get(cur, 0)
                rows.append({"Currency": cur, "Number of Transactions": len(sub),
                             "Total Amount": sub[amount].sum(), "%": len(sub) / base if base else 0})
            rows.append({"Currency": None, "Number of Transactions": len(part), "Total Amount": None,
                         "%": len(part) / len(df) if len(df) else 0})
            out = pd.DataFrame(rows)
            return out if with_pct else out.drop(columns="%")

        return {
            "Overall processed payments": table(df, with_pct=False),
            "Overall notifications received by Irembo": table(rep[m["received"]]),
            "SLA Violated": table(rep[m["sla"]]),
            "Pending notifications": table(rep[m["pending"]]),
            "Duplicate Payment": table(rep[m["duplicate"]]),
        }

    def build_merchants(a, b):
        """Payment per merchant for [a, b): the notifications received (the same rows as that table), one row per
        merchant and currency - amounts are never added across currencies -, most transactions first, then an
        "All" row per currency. % = share of the currency's notifications received."""
        part = rep[table_rows(a, b)["received"]]
        name = part["name"].astype(object) if "name" in part else pd.Series(None, index=part.index, dtype=object)
        name = name.where(name.notna() & name.astype(str).str.strip().ne(""), "(no merchant)").astype(str).str.strip()
        rows = []
        for cur in list(bank.CURRENCIES) + sorted(set(part["Currency"].dropna()) - set(bank.CURRENCIES)):
            sub = part[part["Currency"] == cur]
            if sub.empty:
                continue
            g = (pd.DataFrame({"m": name[sub.index], "amt": sub[amount]}).groupby("m")
                 .agg(n=("m", "size"), amt=("amt", "sum")).reset_index().sort_values(["n", "m"], ascending=[False, True]))
            total = len(sub)
            rows += [{"Merchant name": m, "Currency": cur, "Number of Transactions": int(n), "Total Amount": amt,
                      "%": n / total} for m, n, amt in zip(g["m"], g["n"], g["amt"])]
            rows.append({"Merchant name": "All", "Currency": cur, "Number of Transactions": total,
                         "Total Amount": sub[amount].sum(), "%": 1.0})
        if not rows:
            rows = [{"Merchant name": "All", "Currency": bank.CURRENCIES[0], "Number of Transactions": 0,
                     "Total Amount": 0.0, "%": 0.0}]
        return pd.DataFrame(rows)

    # one block for a day or a week; for a month one block per week (Monday to Sunday) + one for the month.
    # Headings show the last day of the block (the end date itself is not included).
    one_day = pd.Timedelta(days=1)
    if period.report_type == "day":
        res.blocks = [("Summary-Daily", f"Day: {start:%d/%m/%Y}", build_summary5(start, end))]
        res.merchants[res.blocks[-1][1]] = build_merchants(start, end)
    else:
        if period.report_type == "week":
            weeks = [(start, end)]
        else:
            cuts = [start] + [d for d in pd.date_range(start + pd.Timedelta(days=1), end - pd.Timedelta(days=1))
                              if d.dayofweek == 0]
            weeks = list(zip(cuts, cuts[1:] + [end]))
        for i, (a, b) in enumerate(weeks, 1):
            # a weekly report: its week name (Week 1 · October 2026); a monthly report: its blocks in order
            name = week_name(period) if period.report_type == "week" else f"Week {i}"
            res.blocks.append(("Summary-Weekly", f"{name}: {a:%d/%m/%Y} - {b - one_day:%d/%m/%Y}",
                               build_summary5(a, b)))
            res.merchants[res.blocks[-1][1]] = build_merchants(a, b)
        if period.report_type == "month":
            res.blocks.append(("Summary-Monthly", f"Month: {start:%B %Y} ({start:%d/%m/%Y} - {period.last_day:%d/%m/%Y})",
                               build_summary5(start, end)))
            res.merchants[res.blocks[-1][1]] = build_merchants(start, end)

    # the same counts for each day (only used by the charts of a custom report)
    res.daily = {d: counts_of(build_summary5(d, d + one_day)) for d in pd.date_range(start, period.last_day)}

    # ---- G. totals check: the weeks must add up to the month ----
    if period.report_type == "month":
        month = counts_of(res.blocks[-1][2])
        weeks_sum = {}
        for _, _, tables in res.blocks[:-1]:
            for title, d in counts_of(tables).items():
                for cur, n in d.items():
                    weeks_sum.setdefault(title, {}).setdefault(cur, 0)
                    weeks_sum[title][cur] += n
        skip = option(bank, "TOTALS_CHECK_SKIP", [])     # tables whose weeks may overlap by design
        diffs = [f"{title} {cur}: weeks {weeks_sum[title][cur]:,} vs month {n:,}"
                 for title, d in month.items() for cur, n in d.items()
                 if weeks_sum[title][cur] != n and title not in skip]
        if diffs:
            res.add("warning", "Totals check: the weekly blocks do not add up to the month: " + "; ".join(diffs))
        else:
            res.add("success", "Totals check: the weekly blocks add up to the month.")

    # ---- H. comparison with the previous period of the same bank (only when ticked) ----
    if compare:
        prev_path, prev = find_previous_report(base_dir, period, bank)
        if previous_report is not None:
            prev_path = Path(previous_report)
        if prev_path is None:
            raise ReconError(f"No {bank.NAME} report found for {prev.name}. Upload it or make it first.")
        prev_no_tx = is_no_transaction_report(prev_path, prev)   # the bank had no transactions then: all 0
        if prev_no_tx:
            prev_counts, prev_unreviewed = read_previous_counts(prev_path, prev, bank), 0
        else:
            prev_counts, _, prev_unreviewed = read_report_counts(prev_path, SUMMARY_SHEET[period.report_type], bank)
        res.compared_with = (f"{prev.name} (" + ("No-Transaction Report" if prev_no_tx else "report")
                             + f" of {pd.Timestamp(prev_path.stat().st_mtime, unit='s'):%d/%m/%Y})")
        res.previous_counts = prev_counts
        if prev_unreviewed:
            res.add("warning", f"{prev.name} was not reviewed: its {prev_unreviewed} missing invoice(s) are counted as "
                               "not found. The comparison may not be exact.")
        prev_merchants = read_merchant_counts(prev_path, SUMMARY_SHEET[period.report_type])   # {} when it has none
        week_blocks = [b for b in res.blocks if b[0] == "Summary-Weekly"]
        new_blocks = []
        for block in res.blocks:
            sheet_name, heading, tables = block
            if sheet_name == "Summary-Weekly" and period.report_type == "month":
                i = week_blocks.index(block)
                prev_tables = counts_of(week_blocks[i - 1][2]) if i > 0 else None      # week 1: no comparison
                new_blocks.append((sheet_name, heading, add_vs_column(tables, prev_tables, "Week")))
                prev_m = merchant_counts(res.merchants.get(week_blocks[i - 1][1])) if i > 0 else None
                res.merchants[heading] = add_vs_merchants(res.merchants[heading], prev_m, "Week")
            else:
                new_blocks.append((sheet_name, heading, add_vs_column(tables, prev_counts, period.unit, prev_no_tx)))
                res.merchants[heading] = add_vs_merchants(res.merchants[heading], prev_merchants, period.unit,
                                                          prev_no_tx)
        res.blocks = new_blocks
    res.timings["summary"] = time.perf_counter() - t

    # ---- I. report: the detail sheets, the Excel file, the run log ----
    t = time.perf_counter()
    statement_cols = ["Currency"] + present(posted, paid_on, desc, ref, amount)
    detail_cols = statement_cols + option(bank, "DETAIL_COLUMNS", []) + [
        "invoice_number", "found_in_irembopay", "amount", "currency", "payment_status",
        "payment_provider", "payment_channel", "transaction_status", "name", "created_at",
        "paid_at", "settled_at", "provider_reference", "transaction_reference",
        sla_col, "sla_status", "review_status", "review_comment",
        "notification_received", "times_on_statement", "is_extra_credit", "Reason", "Report note", "IremboPay check",
        "Source file"]
    whole = table_rows(start, end)
    rep["Report note"] = table_notes(rep, whole, bank)          # an invoice credited more than once: why, where
    detail = rep[[col for col in detail_cols if col in rep.columns]]
    sheets = {"Pending notification": detail[whole["pending"]].sort_values(paid_on),
              "Duplicate payment other channel": detail[whole["duplicate"]].sort_values(["payment_provider", paid_on]),
              bank.STATEMENT_SHEET: rep[statement_cols + ["invoice_number", "Source file"]]}
    if hasattr(bank, "extra_sheets"):
        sheets.update(bank.extra_sheets(rep, attempts, start, end, whole))
    sheets["Payment details"] = detail
    res.sheets = sheets
    report_file = report_path(base_dir, period, bank, report_dir)
    report_file.parent.mkdir(parents=True, exist_ok=True)
    _write_report(report_file, res.blocks, sheets, bank, res.reconciled_on, res.merchants)
    res.report_file = report_file
    res.timings["report"] = time.perf_counter() - t
    res.add("success", f"Report made: {report_file.name}" if report_dir is not None
            else f"Report saved: output/{period.report_type}/{report_file.name}")
    _log_run(output_dir, res, sla_days)
    return res


# =====================================================================================================================
# 10. RUN LOG (output/run_log.csv: one line per run, shown in "Recent runs" in the side panel)
# =====================================================================================================================
def report_figures(blocks: list) -> dict:
    """The figures of a report, read from its main summary block (the month, or the day / week): the same numbers as
    the Excel, the PDF, Slack and the dashboard (20/10/2026: processed = the invoices on both sides, each once).
    -> {"processed", "received", "rate", "currencies": [(currency, processed, received)]}."""
    main = next((b for b in blocks if b[0] in ("Summary-Monthly", "Summary-Daily")), blocks[-1] if blocks else None)
    if main is None:
        return {"processed": 0, "received": 0, "rate": None, "currencies": []}
    proc, rec = main[2]["Overall processed payments"], main[2]["Overall notifications received by Irembo"]
    total = lambda t: int(t.loc[t["Currency"].isna(), "Number of Transactions"].iloc[0])
    by_cur = lambda t: {c: int(n) for c, n in zip(t["Currency"], t["Number of Transactions"]) if pd.notna(c)}
    p, r = total(proc), total(rec)
    pc, rc = by_cur(proc), by_cur(rec)
    return {"processed": p, "received": r, "rate": r / p if p else None,
            "currencies": [(c, pc[c], rc.get(c, 0)) for c in pc if pc[c] or rc.get(c, 0)]}


def statement_invoices(files: list, bank=None) -> list:
    """Every invoice number of the uploaded statements (found the way the run finds them: BPR also by its references
    in the IremboPay file), for the checkout check made before the run (CHECKOUT_ALL)."""
    bank = get_bank(bank)
    statements = [f.data for f in files if f.kind == "bank"]
    pays = [f.data for f in files if f.kind == "payments"]
    if not statements:
        return []
    rows = bank.clean_statement(pd.concat(statements, ignore_index=True))
    if hasattr(bank, "find_invoices") and pays:
        rows = bank.find_invoices(rows, pd.concat(pays, ignore_index=True))
    return sorted({str(i) for i in rows["invoice_number"].dropna() if str(i).strip()})


def _log_run(output_dir: Path, res: RunResult, sla_days: int):
    """Add one line for this run: when, bank, period, results, and the seconds spent per step."""
    main = next((b for b in res.blocks if b[0] in ("Summary-Monthly", "Summary-Daily")),
                res.blocks[0] if res.blocks else None)
    totals = counts_of(main[2]) if main else {}

    def total(title):
        return totals.get(title, {}).get("Total", "")

    row = {"run_at": pd.Timestamp.now().strftime("%Y-%m-%d %H:%M:%S"),
           "reconciled_on": res.reconciled_on.strftime("%Y-%m-%d"), "bank": res.bank,
           "report_type": res.period.report_type, "period": res.period.text,
           "bank_rows": res.bank_rows, "not_found": res.not_found, "not_reviewed": res.not_reviewed,
           "result": "waiting for review" if res.stopped_for_review else "report made",
           "processed": total("Overall processed payments"),
           "received": total("Overall notifications received by Irembo"),
           "sla_violated": total("SLA Violated"), "pending": total("Pending notifications"),
           "duplicate": total("Duplicate Payment"), "sla_days": sla_days,
           "compared_with": res.compared_with,
           "report_file": res.report_file.name if res.report_file else ""}
    for step in ("read", "match", "review", "summary", "report"):
        row[f"seconds_{step}"] = round(res.timings.get(step, 0.0), 2)
    log = output_dir / "run_log.csv"
    output_dir.mkdir(parents=True, exist_ok=True)
    new = pd.DataFrame([row]).astype(str)
    if log.exists():
        old = pd.read_csv(log, dtype=str)
        if "bank" not in old.columns:                 # runs logged before banks existed are I&M
            old["bank"] = "IM"
        new = pd.concat([old, new], ignore_index=True)
        new = new[list(row) + [c for c in new.columns if c not in row]]
    new.to_csv(log, index=False)


def recent_runs(base_dir: Path, n: int = 8, bank=None) -> pd.DataFrame:
    """The last n runs (of one bank when bank is given), newest first."""
    log = Path(base_dir) / "output" / "run_log.csv"
    if not log.exists():
        return pd.DataFrame()
    try:
        runs = pd.read_csv(log, dtype=str).fillna("")
        if "bank" not in runs.columns:
            runs["bank"] = "IM"
        runs.loc[runs["bank"].eq(""), "bank"] = "IM"
        if bank is not None:
            runs = runs[runs["bank"].eq(get_bank(bank).CODE)]
        return runs.tail(n).iloc[::-1].reset_index(drop=True)
    except Exception:
        return pd.DataFrame()


# =====================================================================================================================
# 11. WRITING THE EXCEL REPORT
#    Summary sheets: each block = a heading + the 5 tables placed like the Google Sheet layout (see GS_SLOTS).
#    Detail sheets: plain tables with a bold header row, filters and number formats.
# =====================================================================================================================
HEADER_FILL = PatternFill("solid", start_color="D9E1F2")
THIN = Side(style="thin", color="000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BOLD = Font(bold=True)
GS_TITLE_FILL = PatternFill("solid", start_color="D0DEE3")
# where each table goes in a block: (column, slot) -> column 1 = A or 6 = F; slot 0/1/2 = rows 3, 11, 19 of the block
GS_SLOTS = {"Overall processed payments": (1, 0), "Overall notifications received by Irembo": (6, 0),
            "SLA Violated": (1, 1), "Pending notifications": (6, 1), "Duplicate Payment": (1, 2)}
MERCHANT_SLOT = (6, 2)                       # Payment per merchant: right of "Duplicate Payment"
GS_STEP = 8                                  # rows between two slots
GS_BLOCK = 2 + 3 * GS_STEP                   # rows taken by one block (heading + 3 slots)
GS_WIDTHS = {"A": 40, "B": 26, "C": 17, "D": 12, "E": 18, "F": 40, "G": 26, "H": 17, "I": 14, "J": 18, "K": 18}
# with a comparison the left tables have one more column (VS prev): the right ones move one column further (G), so an
# empty column (F) still separates them; without a comparison the layout is unchanged (right tables in F)
GS_WIDTHS_COMPARED = {"A": 40, "B": 26, "C": 17, "D": 12, "E": 18, "F": 3, "G": 40, "H": 26, "I": 17, "J": 14,
                      "K": 18, "L": 18}
RECONCILED_CELL = (1, 6)                     # F1 of each summary sheet: "Reconciled on: Thursday 08/10/2026"
RECONCILED_LABEL = "Reconciled on: "


def _style_table_sheet(ws, bank):
    """Format a detail sheet: bold header, frozen first row, filters, column widths, number and date formats."""
    c = get_bank(bank).COLUMNS
    c_amount, c_posted, c_paid = c["amount"], c["posted_date"], c["payment_date"]
    with_time = option(bank, "DATETIME_COLUMNS", [])
    for cell in ws[1]:
        cell.font, cell.fill, cell.border = BOLD, HEADER_FILL, BOX
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col_cells in ws.columns:
        header = col_cells[0].value
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells[:200])
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max(width + 2, 10), 60)
        for cell in col_cells[1:]:
            if header in (c_amount, "amount"):
                cell.number_format = "#,##0.00"
            elif header in with_time:
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            elif header in (c_posted, c_paid):
                cell.number_format = "yyyy-mm-dd"
            elif header in DATE_COLUMNS:
                cell.number_format = "yyyy-mm-dd hh:mm:ss"


def _write_table(ws, top, left, title, table):
    """One summary table: a merged title row, the column names, then one row per currency + the total row."""
    ncols = len(table.columns)
    ws.merge_cells(start_row=top, start_column=left, end_row=top, end_column=left + ncols - 1)
    t = ws.cell(top, left, title)
    t.font, t.alignment = Font(size=14), Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[top].height = 22
    for j, name in enumerate(table.columns):
        ws.cell(top, left + j).fill = GS_TITLE_FILL
        ws.cell(top, left + j).border = BOX
        ws.cell(top + 1, left + j, name).border = BOX
    for i, rec in enumerate(table.itertuples(index=False), start=2):
        for j, value in enumerate(rec):
            c = ws.cell(top + i, left + j, None if pd.isna(value) else value)
            c.border = BOX
            name = table.columns[j]
            if name in ("Number of Transactions", "Total Amount"):
                c.number_format = "#,##0"
            elif name == "%":
                c.number_format = "0.00%"
            elif name.startswith("VS prev"):
                c.alignment = Alignment(horizontal="right")


def _write_summary(ws, blocks, reconciled_on=None, merchants=None):
    """A summary sheet: the day it was reconciled (F1), then the blocks one under the other (heading, the 5 tables,
    and Payment per merchant in the 6th place). A block is GS_BLOCK rows, more when its merchant table is long."""
    compared = any(str(c).startswith("VS prev") for _, tables in blocks for t in tables.values() for c in t.columns)
    shift = 1 if compared else 0                  # the right tables: column F, or G when there is a comparison
    for col, w in (GS_WIDTHS_COMPARED if compared else GS_WIDTHS).items():
        ws.column_dimensions[col].width = w
    if reconciled_on is not None:
        ws.cell(RECONCILED_CELL[0], RECONCILED_CELL[1] + shift,
                f"{RECONCILED_LABEL}{reconciled_on:%A %d/%m/%Y}").font = Font(size=12, italic=True)
    base = 1
    for heading, tables in blocks:
        ws.cell(base, 1, heading).font = Font(size=14, bold=True)
        for title, table in tables.items():
            col, slot = GS_SLOTS[title]
            _write_table(ws, base + 2 + slot * GS_STEP, col + (shift if col > 1 else 0), title, table)
        height = GS_BLOCK
        table = (merchants or {}).get(heading)
        if table is not None:
            col, slot = MERCHANT_SLOT
            top = base + 2 + slot * GS_STEP
            _write_table(ws, top, col + shift, MERCHANT_TITLE, table)
            height = max(GS_BLOCK, top - base + 2 + len(table) + 3)       # a long table pushes the next block down
        base += height


def _open_report(path: Path):
    """Excel writer: a new file, or the existing one kept when the user added sheets of their own."""
    try:
        if path.exists():
            existing = load_workbook(path, read_only=True)
            own = [n for n in existing.sheetnames if n not in REPORT_SHEETS]
            existing.close()
            if own:
                return pd.ExcelWriter(path, engine="openpyxl", mode="a", if_sheet_exists="replace")
        return pd.ExcelWriter(path, engine="openpyxl")
    except PermissionError:
        raise ReconError(f"{path.name} is open in Excel. Close it and run again.") from None


def _write_report(path: Path, blocks, sheets, bank=None, reconciled_on=None, merchants=None):
    """Write the whole report: summary sheets first, then the detail sheets."""
    summary_sheets = []
    for sheet, _, _ in blocks:
        if sheet not in summary_sheets:
            summary_sheets.append(sheet)
    try:
        with _open_report(path) as writer:
            for sheet in summary_sheets:                       # empty first, so the summaries are the first tabs
                pd.DataFrame().to_excel(writer, sheet_name=sheet, index=False)
            for name, df in sheets.items():
                df.to_excel(writer, sheet_name=name, index=False)
                _style_table_sheet(writer.sheets[name], bank)
            for sheet in summary_sheets:
                _write_summary(writer.sheets[sheet], [(h, t) for s, h, t in blocks if s == sheet], reconciled_on,
                               merchants)
            for old in {NOT_FOUND_SHEET} | OLD_SHEETS:             # left over in a report made before
                if old in writer.book.sheetnames and old not in sheets:
                    del writer.book[old]
    except PermissionError:
        raise ReconError(f"{path.name} is open in Excel. Close it and run again.") from None


NO_TX_STATUS = "No Transactions Recorded"
NO_TX_NOTE = ("No transactions were recorded for the selected reporting period. Consequently, no payment matching or "
              "reconciliation was performed for this period.")
STATUS_FILL, STATUS_INK = PatternFill("solid", start_color="FFF3D6"), "7A4E00"      # the app's amber notice


def previous_column(period: Period) -> str:
    """The name of the column with the previous period's counts: 'Previous week' / 'Previous month' / 'Previous day'."""
    return f"Previous {period.unit.lower()}"


def is_no_transaction_report(path: Path, period: Period) -> bool:
    """True when the report is a No-Transaction Report (its summary sheet has a 'Reconciliation status' row)."""
    wb = load_workbook(path, data_only=True)
    try:
        sheet = SUMMARY_SHEET[period.report_type]
        ws = wb[sheet] if sheet in wb.sheetnames else None
        return ws is not None and any(ws.cell(r, 1).value == "Reconciliation status" for r in range(1, ws.max_row + 1))
    finally:
        wb.close()


def read_previous_counts(path: Path, period: Period, bank=None) -> dict:
    """{table title: {currency: n, "Total": n}} of the previous period's report, for the "Previous week" column:
    a standard report (its first summary block), or a No-Transaction Report (its one table)."""
    bank = get_bank(bank)
    sheet = SUMMARY_SHEET[period.report_type]
    wb = load_workbook(path, data_only=True)
    try:
        ws = wb[sheet] if sheet in wb.sheetnames else None
        no_tx = ws is not None and any(ws.cell(r, 1).value == "Reconciliation status" for r in range(1, ws.max_row + 1))
        if no_tx:
            counts = {title: {cur: 0 for cur in bank.CURRENCIES} for title in GS_SLOTS}
            for r in range(1, ws.max_row + 1):
                title, cur, n = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value
                if title in counts and cur in counts[title] and isinstance(n, (int, float)):
                    counts[title][cur] = int(n)
            for d in counts.values():
                d["Total"] = sum(d.values())
            return counts
    finally:
        wb.close()
    return read_report_counts(path, sheet, bank)[0]


def no_transaction_rows(bank=None, previous: dict | None = None) -> list:
    """The rows of a no-transaction Reconciliation Summary (Excel and the app's preview): the 5 categories in the
    standard report's order, at 0 for each currency of the bank; then the previous period's number of transactions
    (read_previous_counts) and the change against it, as the standard reports' 'VS prev' column: "-100.00%";
    "No change" when the previous period had none too; "N/A" when its report was not found."""
    rows = []
    for title in GS_SLOTS:
        for cur in get_bank(bank).CURRENCIES:
            before = None if previous is None else int((previous.get(title) or {}).get(cur, 0))
            change = ("N/A" if before is None else "No change" if before == 0
                      else f"{(0 - before) / before * 100:.2f}%")
            rows.append([title, cur, 0, 0, "N/A" if before is None else before, change])
    return rows


def no_transaction_excel(period: Period, bank=None, made_at=None, previous: dict | None = None) -> bytes:
    """The report of a period in which the bank had no transactions, in the standard report's look (same title
    fill, borders and number formats): one Reconciliation Summary table (every category at 0 for each currency of
    the bank; the last column = the previous period's transactions)
    and the Reconciliation status under it."""
    from openpyxl import Workbook
    bank = get_bank(bank)
    made_at = pd.Timestamp.now() if made_at is None else pd.Timestamp(made_at)
    wb = Workbook()
    ws = wb.active
    ws.title = SUMMARY_SHEET[period.report_type]
    for col, w in {"A": 40, "B": 14, "C": 24, "D": 18, "E": 18, "F": 20}.items():
        ws.column_dimensions[col].width = w
    ws.cell(1, 1, f"{bank.FULL_NAME} – {period.name[0].upper() + period.name[1:]}").font = Font(size=14, bold=True)
    ws.cell(2, 1, f"Reporting period: {period.start:%d/%m/%Y} to {period.last_day:%d/%m/%Y}").font = Font(italic=True)
    ws.cell(2, 3, f"Generated on: {made_at:%A %d/%m/%Y}").font = Font(italic=True)

    # ---- the summary table: title row, column names, one row per category and currency
    cols = ["Category", "Currency", "Number of Transactions", "Total Amount", previous_column(period),
            f"VS prev {period.unit} (%)"]
    top = 4
    ws.merge_cells(start_row=top, start_column=1, end_row=top, end_column=len(cols))
    t = ws.cell(top, 1, "Reconciliation Summary")
    t.font, t.alignment = Font(size=14), Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[top].height = 22
    for j, name in enumerate(cols, 1):
        ws.cell(top, j).fill, ws.cell(top, j).border = GS_TITLE_FILL, BOX
        cell = ws.cell(top + 1, j, name)
        cell.font, cell.border = BOLD, BOX
    r = top + 2
    for values in no_transaction_rows(bank, previous):
        for j, v in enumerate(values, 1):
            cell = ws.cell(r, j, v)
            cell.border = BOX
            if j in (3, 4) or (j == 5 and isinstance(v, int)):
                cell.number_format = "#,##0"
            elif j in (5, 6):
                cell.alignment = Alignment(horizontal="right")
        r += 1

    # ---- the reconciliation status, under the table
    r += 1
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=len(cols))
    ws.cell(r, 1, "Reconciliation status").font = BOLD
    for j in range(1, len(cols) + 1):
        ws.cell(r, j).fill, ws.cell(r, j).border = GS_TITLE_FILL, BOX
    for i, (text, font) in enumerate([(NO_TX_STATUS, Font(bold=True, size=12, color=STATUS_INK)),
                                      (NO_TX_NOTE, Font(color=STATUS_INK))], 1):
        ws.merge_cells(start_row=r + i, start_column=1, end_row=r + i, end_column=len(cols))
        cell = ws.cell(r + i, 1, text)
        cell.font, cell.alignment = font, Alignment(wrap_text=True, vertical="center", indent=1)
        for j in range(1, len(cols) + 1):
            ws.cell(r + i, j).fill, ws.cell(r + i, j).border = STATUS_FILL, BOX
    ws.row_dimensions[r + 1].height, ws.row_dimensions[r + 2].height = 22, 34
    ws.page_setup.orientation, ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = "portrait", 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# =====================================================================================================================
# 12. SAVED REPORTS AND DECISIONS (the side panel: list, view, delete)
# =====================================================================================================================
def rename_old_reports(base_dir: Path):
    """One time: reports named up to the day AFTER the period ('..._01-09-2026_to_01-10-2026.xlsx') are renamed to
    end on the last day ('..._01-09-2026_to_30-09-2026.xlsx'). A name already taken is left as it is."""
    for kind in SUMMARY_SHEET:
        for p in (Path(base_dir) / "output" / kind).glob(f"*{REPORT_SUFFIX}*_to_*.xlsx"):
            try:
                a, b = (pd.to_datetime(x, format="%d-%m-%Y") for x in report_label(p).split("_to_"))
            except ValueError:
                continue
            period = make_period(kind, a)
            if p.name.startswith("~$") or period.start != a or b != period.end:     # not an old name
                continue
            target = p.with_name(f"{report_bank(p)}{REPORT_SUFFIX}{period.label}.xlsx")
            if not target.exists():
                try:
                    p.rename(target)
                except OSError:
                    pass


def saved_reports(base_dir: Path) -> list:
    """[(kind, path)] of every bank (old report names are renamed first)."""
    rename_old_reports(base_dir)
    out = []
    for kind in ("month", "week", "day"):
        for p in sorted((Path(base_dir) / "output" / kind).glob(f"*{REPORT_SUFFIX}*.xlsx"), reverse=True):
            if not p.name.startswith("~$"):                    # "~$..." = Excel's lock file of an open report
                out.append((kind, p))
    return out


def period_title(kind: str, label: str) -> str:
    """'month', '01-09-2026_to_30-09-2026' -> 'Month · September 2026'. Only the first day is read, so the labels
    that end on the day after the period (review decisions, old reports and run log lines) give the same title."""
    try:
        a = pd.to_datetime(label.split("_to_")[0], format="%d-%m-%Y")
        last = make_period(kind, a).last_day
        if kind == "month":
            return f"Month · {a:%B %Y}"
        if kind == "week":
            month, n = week_number(a)
            return f"Week {n} · {month:%B %Y} ({a:%d/%m} – {last:%d/%m/%Y})"
        return f"Day · {a:%d/%m/%Y}"
    except Exception:
        return f"{kind} · {label}"


def report_title(path: Path, with_bank=True) -> str:
    """'I&M · Month · September 2026' (or 'Month · September 2026')."""
    path = Path(path)
    kind = path.parent.name if path.parent.name in SUMMARY_SHEET else "month"
    title = period_title(kind, report_label(path))
    return f"{bank_name(report_bank(path))} · {title}" if with_bank else title


def report_list(base_dir: Path, bank=None) -> list:
    """[(title, path)] newest period first: every bank, or only one bank (then the title has no bank name)."""
    code = None if bank is None else get_bank(bank).CODE
    items = []
    for kind, path in saved_reports(base_dir):
        if code is not None and report_bank(path) != code:
            continue
        try:
            start = pd.to_datetime(report_label(path).split("_to_")[0], format="%d-%m-%Y")
        except Exception:
            start = pd.Timestamp(0)
        items.append((start, kind, report_title(path, with_bank=code is None), path))
    items.sort(key=lambda x: (x[0], x[1], x[2]), reverse=True)
    return [(title, path) for _, _, title, path in items]


def delete_report(base_dir: Path, path: Path):
    """Delete one report file (only files inside output/ can be deleted)."""
    path = Path(path).resolve()
    output = (Path(base_dir) / "output").resolve()
    if output not in path.parents or path.suffix != ".xlsx":
        raise ReconError("Only reports made by this app can be deleted here.")
    try:
        path.unlink(missing_ok=True)
    except PermissionError:
        raise ReconError(f"{path.name} is open in Excel. Close it and try again.") from None


def review_list(base_dir: Path, bank=None) -> list:
    """[(title, bank code, store_key, number of decisions, last saved)] newest period first: every bank, or only
    one bank (then the title has no bank name)."""
    store = load_store(base_dir, bank)
    out = []
    for (code, key), grp in store.groupby(["bank", "period"]):
        kind, _, label = key.partition(":")
        try:
            start = pd.to_datetime(label.split("_to_")[0], format="%d-%m-%Y")
        except Exception:
            start = pd.Timestamp(0)
        title = period_title(kind, label) if bank is not None else f"{bank_name(code)} · {period_title(kind, label)}"
        out.append((start, title, code, key, len(grp),
                    grp["saved_at"].max()))
    out.sort(key=lambda x: (x[0], x[1]), reverse=True)
    return [x[1:] for x in out]


def delete_review(base_dir: Path, store_key: str, bank=None):
    """Delete the decisions of one period of one bank (those invoices will need a new review)."""
    code = get_bank(bank).CODE
    store = load_store(base_dir)
    _write_store(base_dir, store[~(store["bank"].eq(code) & store["period"].eq(store_key))], "try again")


def _read_merchant_table(ws, top: int, col: int) -> pd.DataFrame:
    """The Payment per merchant table whose title is at (top, col): its header row, then rows down to a blank."""
    headers, c = [], col
    while ws.cell(top + 1, c).value:
        headers.append(ws.cell(top + 1, c).value)
        c += 1
    rows, r = [], top + 2
    while ws.cell(r, col).value not in (None, ""):
        rows.append([ws.cell(r, col + j).value for j in range(len(headers))])
        r += 1
    t = pd.DataFrame(rows, columns=headers)
    if len(t):
        t["Number of Transactions"] = pd.to_numeric(t["Number of Transactions"], errors="coerce").fillna(0).astype(int)
        t["Total Amount"] = pd.to_numeric(t["Total Amount"], errors="coerce")
        t["%"] = pd.to_numeric(t["%"], errors="coerce").fillna(0.0)
    return t


def _right_col(ws, row: int, col: int, title: str):
    """The column of a table titled `title` on `row`: `col`, or the next one (right tables of a compared report)."""
    for c in ((col, col + 1) if col > 1 else (col,)):
        if ws.cell(row, c).value == title:
            return c
    return None


def read_merchant_counts(path: Path, sheet: str) -> dict:
    """{(merchant, currency): transactions} of the first Payment per merchant table of a summary sheet; {} when
    the report has none (made before 12/10/2026): the comparison then says N/A."""
    wb = load_workbook(path, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            return {}
        ws = wb[sheet]
        for r in range(1, ws.max_row + 1):
            col = _right_col(ws, r, MERCHANT_SLOT[0], MERCHANT_TITLE)
            if col is not None:
                return merchant_counts(_read_merchant_table(ws, r, col))
        return {}
    finally:
        wb.close()


class _Cell:
    __slots__ = ("value",)

    def __init__(self, value):
        self.value = value


class _Grid:
    """A summary sheet read once in read-only mode, kept in memory: ws.cell(row, col).value and ws.max_row as with
    openpyxl, without loading the whole workbook (the Payment details sheet is large). Same values, much faster."""

    def __init__(self, ws):
        self.rows = [list(r) for r in ws.iter_rows(values_only=True)]
        self.max_row = len(self.rows)

    def cell(self, row: int, column: int) -> _Cell:
        if 1 <= row <= len(self.rows) and 1 <= column <= len(self.rows[row - 1]):
            return _Cell(self.rows[row - 1][column - 1])
        return _Cell(None)


def read_report(path: Path, count_reviewed: bool = True) -> dict:
    """Read a saved report back: {"report_type", "bank", "title", "blocks": [(sheet, heading, tables)], "not_found",
    "reconciled_on" ('Thursday 08/10/2026', or "" for reports made before it was recorded),
    "merchants": {heading: Payment per merchant table} (empty for reports made before 12/10/2026)}.
    Works for reports made by this app (same layout as _write_summary). count_reviewed=False skips the count of
    reviewed rows ("not_found" = 0) and so the large Payment details sheet: the dashboard does not use it."""
    path = Path(path)
    kind = path.parent.name if path.parent.name in SUMMARY_SHEET else "month"
    code = report_bank(path)
    currencies = BANKS[code].CURRENCIES if code in BANKS else []
    wb = load_workbook(path, data_only=True, read_only=True)     # read only: the summary sheets are copied to memory
    try:
        blocks, reconciled_on, merchants = [], "", {}
        for sheet in ("Summary-Daily", "Summary-Weekly", "Summary-Monthly"):
            if sheet not in wb.sheetnames:
                continue
            ws = _Grid(wb[sheet])
            stamp = str(ws.cell(*RECONCILED_CELL).value or "") or str(ws.cell(RECONCILED_CELL[0],
                                                                                  RECONCILED_CELL[1] + 1).value or "")
            if stamp.startswith(RECONCILED_LABEL):
                reconciled_on = stamp[len(RECONCILED_LABEL):]
            # a block = a heading in column A with "Overall processed payments" two rows below (found by looking:
            # a block is taller when its merchant table is long; reports made before have the same layout)
            starts = [r for r in range(1, ws.max_row + 1)
                      if ws.cell(r, 1).value and ws.cell(r + 2, 1).value == "Overall processed payments"]
            for base in starts:
                heading = ws.cell(base, 1).value
                tables = {}
                for title, (col, slot) in GS_SLOTS.items():
                    top = base + 2 + slot * GS_STEP
                    col = _right_col(ws, top, col, title)
                    if col is None:
                        continue
                    headers = []
                    c = col
                    while (c < col + 5 and ws.cell(top + 1, c).value       # a table has 5 columns at most: stop at
                           and (c == col or ws.cell(top + 1, c).value != "Currency")):   # the next table's columns
                        headers.append(ws.cell(top + 1, c).value)
                        c += 1
                    n_rows = len(currencies) + 1
                    if not currencies:                                    # unknown bank: rows down to "Total"
                        while ws.cell(top + 1 + n_rows, col).value not in (None, "") and n_rows < 20:
                            n_rows += 1
                    rows = [[ws.cell(top + 2 + i, col + j).value for j in range(len(headers))]
                            for i in range(n_rows)]
                    t = pd.DataFrame(rows, columns=headers)
                    t["Number of Transactions"] = pd.to_numeric(t["Number of Transactions"], errors="coerce").fillna(0).astype(int)
                    t["Total Amount"] = pd.to_numeric(t["Total Amount"], errors="coerce")
                    if "%" in t:
                        t["%"] = pd.to_numeric(t["%"], errors="coerce").fillna(0.0)
                    tables[title] = t
                if len(tables) == len(GS_SLOTS):
                    blocks.append((sheet, heading, tables))
                    col, slot = MERCHANT_SLOT
                    col = _right_col(ws, base + 2 + slot * GS_STEP, col, MERCHANT_TITLE)
                    if col is not None:
                        merchants[heading] = _read_merchant_table(ws, base + 2 + slot * GS_STEP, col)
        not_found = 0                                             # rows that went through the review
        if count_reviewed and "Payment details" in wb.sheetnames:
            ws = wb["Payment details"]
            headers = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ()))
            if "review_status" in headers:
                col = headers.index("review_status")
                not_found = sum(1 for row in ws.iter_rows(min_row=2, values_only=True) if len(row) > col and row[col])
    finally:
        wb.close()
    if not blocks:
        raise ReconError(f"{path.name}: this report cannot be shown (it was not made by this app).")
    return {"report_type": kind, "bank": bank_name(code), "full_name": BANKS[code].FULL_NAME if code in BANKS else code,
            "title": period_title(kind, report_label(path)), "blocks": blocks, "not_found": not_found,
            "reconciled_on": reconciled_on, "merchants": merchants}


def _notebook_table(ws, top: int, n_currencies: int, col: int = 2) -> pd.DataFrame:
    """One table of a notebook report's Summary sheet, starting at its own title's column (B for Overall processed
    payments / SLA Violated / Duplicate Payment - one column to the right of where this app puts its own, see
    report_history.parse_notebook_name; further right, at whatever column, for Overall notifications received by
    Irembo / Pending notifications, when the notebooks wrote them - found by read_notebook_report, not assumed
    here)."""
    headers, c = [], col
    while c < col + 5 and ws.cell(top + 1, c).value and (c == col or ws.cell(top + 1, c).value != "Currency"):
        headers.append(ws.cell(top + 1, c).value)
        c += 1
    rows = [[ws.cell(top + 2 + i, col + j).value for j in range(len(headers))] for i in range(n_currencies + 1)]
    t = pd.DataFrame(rows, columns=headers)
    t["Number of Transactions"] = pd.to_numeric(t["Number of Transactions"], errors="coerce").fillna(0).astype(int)
    t["Total Amount"] = pd.to_numeric(t["Total Amount"], errors="coerce")
    return t


def _notebook_block_starts(ws) -> list:
    """[(row, col), ...] of every "Overall processed payments" table in a notebook Summary sheet - one per week
    block (Summary-Weekly: a week each) or the one block of Summary-Monthly, wherever its own heading is (above
    it, two rows up, in the same column)."""
    out = []
    for r in range(1, ws.max_row + 1):
        for c in range(1, 12):
            if ws.cell(r, c).value and ws.cell(r + 2, c).value == "Overall processed payments":
                out.append((r, c))
                break
    return out


_NB_BLOCK_DATES = re.compile(r"\((\d{4}-\d{2}-\d{2})\s+to\s+(\d{4}-\d{2}-\d{2})\)")


def _notebook_block_tables(ws, row: int, col: int, bank, pending_fallback: pd.DataFrame | None) -> dict:
    """The 5 summary tables of one block: Overall processed payments / SLA Violated / Duplicate Payment read from
    the Summary sheet; Pending notifications read from the same row group when the notebooks wrote it there (any
    column, searched for - not assumed to be at a fixed offset), else `pending_fallback` (None: counted as 0, not
    a guess - used for a week block, whose own pending cannot be told apart in the notebooks' one "Pending
    notification" sheet for the whole report). Overall notifications received by Irembo is always derived
    (Processed - Pending - Duplicate), never read: the notebooks' own column under that title can mirror Processed
    instead of Notified (CLAUDE.md; found 2026-10-28, not trusted here)."""
    n = len(bank.CURRENCIES)
    processed = _notebook_table(ws, row + 2, n, col)
    sla_row = row + 2 + GS_STEP
    sla = _notebook_table(ws, sla_row, n, col)
    duplicate = _notebook_table(ws, row + 2 + 2 * GS_STEP, n, col)
    pending = None
    for c in range(1, 12):
        if ws.cell(sla_row, c).value == "Pending notifications":
            pending = _notebook_table(ws, sla_row, n, c)
            break
    if pending is None:
        pending = (pending_fallback if pending_fallback is not None
                  else _notebook_difference(processed, processed, currencies=bank.CURRENCIES))
    notified = _notebook_difference(processed, pending, duplicate, currencies=bank.CURRENCIES)
    return {"Overall processed payments": processed, "Overall notifications received by Irembo": notified,
           "SLA Violated": sla, "Pending notifications": pending, "Duplicate Payment": duplicate}


def _notebook_sheet_counts(wb, sheet_name: str, currencies: list) -> pd.DataFrame | None:
    """A table in the same shape as _notebook_table, counted from a notebook report's per-invoice sheet (Pending
    notification, Duplicate payment other channel) instead of read from the Summary - the notebooks never write
    these two as a Summary table. None when the sheet is not there at all."""
    if sheet_name not in wb.sheetnames:
        return None
    ws = wb[sheet_name]
    headers = [c.value for c in ws[1]]
    if "Ccy" not in headers or "Amount" not in headers:
        return None
    ci, ai = headers.index("Ccy"), headers.index("Amount")
    rows = []
    for cur in currencies:
        n, amount = 0, 0.0
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, ci + 1).value == cur:
                n += 1
                amount += float(ws.cell(r, ai + 1).value or 0)
        rows.append({"Currency": cur, "Number of Transactions": n, "Total Amount": amount})
    rows.append({"Currency": None, "Number of Transactions": sum(r["Number of Transactions"] for r in rows),
                "Total Amount": sum(r["Total Amount"] for r in rows)})
    return pd.DataFrame(rows)


def _notebook_difference(minuend: pd.DataFrame, *subtrahends, currencies: list) -> pd.DataFrame:
    """minuend minus each of subtrahends, currency by currency (count and amount) - a missing subtrahend (no
    Pending notification sheet, say) counts as 0. Used for the two tables the notebooks never write at all
    (Overall notifications received by Irembo = Processed - Pending - Duplicate): for Ecobank this can come out
    surprising, even negative - the same gap already known for its Processed figure (CLAUDE.md), not something
    introduced here."""
    def by_currency(t):
        if t is None:
            return {}
        return {r["Currency"]: (r["Number of Transactions"], 0.0 if pd.isna(r["Total Amount"]) else r["Total Amount"])
               for _, r in t.iterrows() if pd.notna(r["Currency"])}
    base = by_currency(minuend)
    subs = [by_currency(t) for t in subtrahends]
    rows = []
    for cur in currencies:
        n, amount = base.get(cur, (0, 0.0))
        for s in subs:
            sn, samount = s.get(cur, (0, 0.0))
            n, amount = n - sn, amount - samount
        rows.append({"Currency": cur, "Number of Transactions": n, "Total Amount": amount})
    rows.append({"Currency": None, "Number of Transactions": sum(r["Number of Transactions"] for r in rows),
                "Total Amount": sum(r["Total Amount"] for r in rows)})
    return pd.DataFrame(rows)


def read_notebook_report(path: Path, bank, period: Period) -> dict:
    """Read a report made by the reconciliation notebooks (CLAUDE.md's reference notebooks), not this app: same
    return shape as read_report(), so the Dashboard and Report History do not need to know which kind they got.
    A monthly report also has its own Summary-Weekly sheet, one block per week of that month (the notebooks write
    it next to Summary-Monthly) - both are read, so the Dashboard's per-week charts have real weeks instead of
    none at all (2026-10-28). No merchant table, no "reconciled on" stamp, no "needs a look" count - the notebooks
    do not keep those."""
    bank = get_bank(bank)
    sheet = SUMMARY_SHEET[period.report_type]
    wb = load_workbook(path, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise ReconError(f"{path.name}: not a report the notebooks made (no {sheet} sheet).")
        pending_sheet = _notebook_sheet_counts(wb, "Pending notification", bank.CURRENCIES)
        ws = wb[sheet]
        starts = _notebook_block_starts(ws)
        if not starts:
            raise ReconError(f"{path.name}: not a report the notebooks made (no summary table found).")
        row, col = starts[0]                        # the period's own total: the month, or the lone week
        if period.report_type == "month":
            heading = f"Month: {period.start:%B %Y} ({period.start:%d/%m/%Y} - {period.last_day:%d/%m/%Y})"
        else:
            heading = f"Week 1: {period.start:%d/%m/%Y} - {period.last_day:%d/%m/%Y}"
        blocks = [(sheet, heading, _notebook_block_tables(ws, row, col, bank, pending_sheet))]
        if period.report_type == "month" and "Summary-Weekly" in wb.sheetnames:
            wws = wb["Summary-Weekly"]
            for r, c in _notebook_block_starts(wws):
                dates = _NB_BLOCK_DATES.search(str(wws.cell(r, c).value or ""))
                if not dates:                       # a week heading not in the expected format: left out
                    continue
                a, b = (pd.Timestamp(x) for x in dates.groups())
                blocks.append(("Summary-Weekly", f"Week: {a:%d/%m/%Y} - {b:%d/%m/%Y}",
                              _notebook_block_tables(wws, r, c, bank, None)))
    finally:
        wb.close()
    label = f"{period.start:%d-%m-%Y}_to_{period.last_day:%d-%m-%Y}"
    return {"report_type": period.report_type, "bank": bank_name(bank.CODE), "full_name": bank.FULL_NAME,
            "title": period_title(period.report_type, label), "blocks": blocks, "not_found": 0,
            "reconciled_on": "", "merchants": {}}


# =====================================================================================================================
# 13. PDF (summary only): page 1 = key figures + the main summary; then one page per week
# =====================================================================================================================
def report_pdf(path: Path) -> bytes:
    """Summary-only PDF of a saved report: key figures + main summary on page 1, then one page per week."""
    from reportlab.graphics.charts.barcharts import VerticalBarChart
    from reportlab.graphics.shapes import Drawing
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import KeepTogether, LongTable, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape

    rep = read_report(path)
    teal, green = colors.HexColor("#123239"), colors.HexColor("#99E378")
    title_fill, grid = colors.HexColor("#D0DEE3"), colors.HexColor("#9AA8AD")
    warning_fill, warning_stroke = colors.HexColor("#FFF3D6"), colors.HexColor("#D4A72C")
    muted = colors.HexColor("#5B6B70")
    made = pd.Timestamp(Path(path).stat().st_mtime, unit="s")
    page_w, page_h = landscape(A4)
    margin = 14 * mm
    usable = page_w - 2 * margin
    h1 = ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=17, textColor=green, leading=22)
    sub = ParagraphStyle("sub", fontName="Helvetica", fontSize=9.5, textColor=colors.HexColor("#D6E4E6"), leading=12)
    h2 = ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=12.5, textColor=teal, leading=16, spaceAfter=4)
    h3 = ParagraphStyle("h3", fontName="Helvetica-Bold", fontSize=9.2, textColor=teal, leading=12, spaceAfter=3)
    small = ParagraphStyle("small", fontName="Helvetica", fontSize=8.1, textColor=muted, leading=10)
    cell = ParagraphStyle("cell", fontName="Helvetica", fontSize=8.2, leading=10)

    def title_band(subtitle, warning=False):
        """Balanced title band with the current report metadata."""
        label = (rep.get("full_name") or "Bank Reconciliation").strip()
        t = Table([[Paragraph(f"{escape(label)} Reconciliation", h1)], [Paragraph(subtitle, sub)]], colWidths=[usable])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), teal),
            ("LEFTPADDING", (0, 0), (-1, -1), 12),
            ("RIGHTPADDING", (0, 0), (-1, -1), 12),
            ("TOPPADDING", (0, 0), (-1, -1), 8),
            ("BOTTOMPADDING", (0, -1), (-1, -1), 8),
            ("LINEBEFORE", (0, 0), (0, -1), 5, green if not warning else warning_stroke),
            ("TEXTCOLOR", (0, 0), (-1, -1), colors.white),
        ]))
        return t

    def fmt_value(v):
        if v is None or (not isinstance(v, str) and pd.isna(v)):
            return ""
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            if abs(v) >= 1000:
                return f"{v:,.0f}"
            return f"{v:,.2f}" if isinstance(v, float) and not v.is_integer() else f"{v:,.0f}"
        return str(v)

    def one_table(title, t, width, warning=False):
        """One summary table with consistent numeric alignment and a clear total row."""
        cols = list(t.columns)
        head = {"Number of Transactions": "Transactions", "Total Amount": "Total amount", "%": "% of processed"}
        rows = [[title] + [""] * (len(cols) - 1), [head.get(c, c) for c in cols]]
        for rec in t.itertuples(index=False):
            out = []
            for c, v in zip(cols, rec):
                if c == "Currency":
                    out.append("Total" if pd.isna(v) else v)
                elif c == "Number of Transactions":
                    out.append(f"{int(v):,}")
                elif c == "Total Amount":
                    out.append("" if pd.isna(v) else f"{v:,.0f}")
                elif c == "%":
                    out.append(f"{v:.2%}")
                else:
                    out.append("" if v is None or (isinstance(v, float) and pd.isna(v)) else str(v))
            rows.append(out)
        first = 0.24 * width
        widths = [first] + [(width - first) / (len(cols) - 1)] * (len(cols) - 1)
        tb = LongTable(rows, colWidths=widths, repeatRows=2)
        title_bg = warning_fill if warning else title_fill
        title_line = warning_stroke if warning else grid
        style = [
            ("SPAN", (0, 0), (-1, 0)),
            ("BACKGROUND", (0, 0), (-1, 0), title_bg),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 9.2),
            ("TEXTCOLOR", (0, 0), (-1, 0), teal if not warning else colors.HexColor("#563B00")),
            ("ALIGN", (0, 0), (-1, 0), "CENTER"),
            ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 1), (-1, -1), 8.2),
            ("FONTNAME", (-1, 2), (-1, -1), "Helvetica-Bold"),
            ("ALIGN", (1, 2), (-1, -1), "RIGHT"),
            ("GRID", (0, 0), (-1, -1), 0.45, title_line),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("ROWBACKGROUNDS", (0, 2), (-1, -1), [colors.white, colors.HexColor("#F7FAFB")]),
        ]
        if warning:
            style.extend([
                ("LINEABOVE", (0, 1), (-1, 1), 1.4, warning_stroke),
                ("TEXTCOLOR", (0, 2), (-1, -1), colors.HexColor("#563B00")),
            ])
        tb.setStyle(TableStyle(style))
        return tb

    def summary_grid(tables):
        """The 5 summary tables, arranged in a compact 2-column grid."""
        half = (usable - 8 * mm) / 2
        order = [["Overall processed payments", "Overall notifications received by Irembo"],
                 ["SLA Violated", "Pending notifications"], ["Duplicate Payment", None]]
        rows = [[one_table(a, tables[a], half, warning=(a == "SLA Violated")),
                 one_table(b, tables[b], half, warning=(b == "SLA Violated")) if b else ""]
                for a, b in order]
        g = Table(rows, colWidths=[half + 4 * mm, half + 4 * mm])
        g.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
            ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ("RIGHTPADDING", (0, 0), (-1, -1), 8 * mm),
        ]))
        return g

    def distribution_chart(tables, width, height):
        """Small compact distribution chart for received, pending, duplicate and not-found totals."""
        counts = counts_of(tables)
        labels = ["Received notification", "Pending", "Duplicate", "Not found"]
        values = [
            counts.get("Overall notifications received by Irembo", {}).get("Total", 0),
            counts.get("Pending notifications", {}).get("Total", 0),
            counts.get("Duplicate Payment", {}).get("Total", 0),
            rep.get("not_found", 0),
        ]
        if sum(values) <= 0:
            return None
        d = Drawing(width, height)
        bc = VerticalBarChart()
        bc.x = 22
        bc.y = 16
        bc.height = max(52, height - 30)
        bc.width = width - 32
        bc.data = [values]
        bc.categoryAxis.categoryNames = labels
        bc.categoryAxis.labels.fontName = "Helvetica"
        bc.categoryAxis.labels.fontSize = 7
        bc.categoryAxis.strokeColor = grid
        bc.categoryAxis.visibleTicks = False
        bc.valueAxis.valueMin = 0
        bc.valueAxis.visibleGrid = True
        bc.valueAxis.gridStrokeColor = colors.HexColor("#E3ECEF")
        bc.valueAxis.strokeColor = grid
        bc.valueAxis.labels.fontName = "Helvetica"
        bc.valueAxis.labels.fontSize = 7
        bc.valueAxis.labelTextFormat = lambda v: f"{int(v):,}"
        bc.bars[0].fillColor = teal
        bc.bars[0].strokeColor = colors.HexColor("#122E36")
        bc.groupSpacing = 18
        d.add(bc)
        return d

    def kpis(tables, not_found):
        """The six key figures, styled as a compact KPI grid."""
        def total(title):
            t = tables[title]
            row = t[t["Currency"].isna()]
            return int(row["Number of Transactions"].iloc[0]), (row["%"].iloc[0] if "%" in t else None)

        processed, _ = total("Overall processed payments")
        rec_n, rec_pct = total("Overall notifications received by Irembo")
        sla, _ = total("SLA Violated")
        values = [
            ("Processed transaction", f"{processed:,}"),
            ("Received notification", f"{rec_n:,} ({rec_pct:.2%})"),
            ("SLA violated", f"{sla:,}"),
            ("Pending (NEW)", f"{total('Pending notifications')[0]:,}"),
            ("Duplicate", f"{total('Duplicate Payment')[0]:,}"),
            ("Not found", f"{not_found:,}"),
        ]

        cards = []
        for label, value in values:
            cards.append(Table([[Paragraph(escape(label), small)], [Paragraph(escape(value), ParagraphStyle("kpi", fontName="Helvetica-Bold", fontSize=14, textColor=teal, leading=18))]], colWidths=[usable / 6]))
            cards[-1].setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), colors.white),
                ("BOX", (0, 0), (-1, -1), 0.6, grid),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]))
        row = Table([cards], colWidths=[usable / 6] * 6)
        row.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
        return row

    main = next((b for b in rep["blocks"] if b[0] in ("Summary-Monthly", "Summary-Daily")), rep["blocks"][0])
    weeks = [b for b in rep["blocks"] if b[0] == "Summary-Weekly" and b is not main]
    when = f"reconciled on {rep['reconciled_on']}" if rep["reconciled_on"] else f"made on {made:%d/%m/%Y %H:%M}"
    report_title = rep.get("title") or "Report"
    subtitle = f"{report_title} · {when}"

    story = [title_band(subtitle)]
    story += [Spacer(1, 5 * mm), kpis(main[2], rep["not_found"])]

    chart = distribution_chart(main[2], usable * 0.52, 70 * mm)
    if chart is not None:
        chart_block = Table([[Paragraph("Transaction distribution", h3), ""], [chart, ""]], colWidths=[usable * 0.52, usable * 0.48])
        chart_block.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        story += [Spacer(1, 5 * mm), chart_block]

    story += [Spacer(1, 5 * mm), Paragraph(main[1], h2), KeepTogether(summary_grid(main[2]))]

    for _, heading, tables in weeks:
        story += [PageBreak(), title_band(f"{report_title} · {heading.split(':')[0]}", warning=heading.startswith("Week") and "SLA" in str(tables))]
        story += [Spacer(1, 5 * mm), Paragraph(heading, h2), KeepTogether(summary_grid(tables))]

    def footer(canvas, doc):
        """Page number and report identity on every page."""
        canvas.saveState()
        canvas.setFont("Helvetica-Bold", 8)
        canvas.setFillColor(teal)
        canvas.drawString(margin, 8 * mm, f"{report_title} · {rep.get('full_name') or 'Report'}")
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(muted)
        canvas.drawRightString(page_w - margin, 8 * mm, f"Page {doc.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=margin, rightMargin=margin,
                            topMargin=margin, bottomMargin=margin, title=f"{rep.get('full_name') or 'Reconciliation'} {report_title}")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


# =====================================================================================================================
# 14. CUSTOM REPORT (step 6, "Custom report": the user picks the sections, how each is shown, and the layout)
#    custom_content() builds the content once from the run; the app's preview, custom_pdf() and custom_excel() all
#    draw that same content, so they always match. The standard report in output/ is never changed.
#    Charts are drawn in the PDF, the preview and the Excel file (real Excel charts, with their numbers beside).
# =====================================================================================================================
TABLE, BAR, LINE, KPI = "Table", "Bar chart", "Line chart", "KPI summary"
CHART_TEAL = "#123239"           # bars and lines of the PDF charts (IremboPay teal: 13.6:1 contrast on white paper)
# key: (title, the summary table it is read from or None, the ways it can be shown: the first is the default)
CUSTOM_SECTIONS = {
    "overview": ("Reconciliation Overview", None, [KPI]),
    "payments": ("Payment Summary", "Overall processed payments", [TABLE, BAR, LINE]),
    "received": ("Notifications Received", "Overall notifications received by Irembo", [TABLE, BAR, LINE]),
    "sla": ("SLA Violations", "SLA Violated", [TABLE, BAR, LINE]),
    "pending": ("Pending Notifications", "Pending notifications", [TABLE, BAR, LINE]),
    "duplicates": ("Duplicate Payments", None, [TABLE]),                       # credited twice on the statement
    "other_channel": ("Other Channel Payments", "Duplicate Payment", [TABLE, BAR, LINE]),   # paid another way
    "merchants": ("Payment per Merchant", None, [TABLE, BAR]),       # notifications received, by merchant
}
# report type: (description, {section: how it is shown}) - choosing a type fills in the sections and views
REPORT_TYPES = {
    "Full Report": ("Detailed reconciliation",
                    {k: v[2][0] for k, v in CUSTOM_SECTIONS.items() if k != "other_channel"}),
    "Management Summary": ("KPIs & trends", {"overview": KPI, "payments": TABLE, "received": BAR, "sla": BAR,
                                             "pending": TABLE}),
    "Pending & Duplicates Report": ("Pending, duplicate & other-channel payments",
                                    {"pending": TABLE, "duplicates": TABLE, "other_channel": TABLE}),
}
CUSTOM_INCLUDE = {"kpis": "Summary KPIs", "currency": "Currency breakdown", "timestamp": "Reconciliation timestamp"}
# once the report is made, the reviewed rows are simply in the table their decision chose: the custom report does not
# list or count them again (the decisions are seen in Exception Resolution only)


def chart_groups(period: Period) -> list:
    """How a chart can group the transactions: by week (a month), by day (a week or a month), by currency."""
    return ((["Week"] if period.report_type == "month" else []) + (["Day"] if period.report_type != "day" else [])
            + ["Currency"])


def _summary_frame(t: pd.DataFrame) -> pd.DataFrame:
    """A summary table for the reader: Currency (Total), Transactions, Total amount, % of processed."""
    out = pd.DataFrame({"Currency": t["Currency"].fillna("Total").values,
                        "Transactions": t["Number of Transactions"].astype(int).values,
                        "Total amount": t["Total Amount"].values})
    if "%" in t:
        out["% of processed"] = t["%"].values
    return out


def _list_frame(df: pd.DataFrame, cols: dict) -> pd.DataFrame:
    """Some columns of a detail sheet, renamed for the reader ({column: name}); missing columns are skipped."""
    keep = {c: n for c, n in cols.items() if c in df.columns}
    return df[list(keep)].rename(columns=keep).reset_index(drop=True)


def _series(res: RunResult, main, weeks, table_title: str, view: dict):
    """([(label, transactions)], group used) for a chart: by week, by day or by currency."""
    col = "Total" if view["currency"] == "All" else view["currency"]
    group = view["group"]
    if group == "Week" and weeks:
        return [(f"Week {i}", counts_of(t)[table_title].get(col, 0)) for i, (_, _, t) in enumerate(weeks, 1)], group
    if group == "Day" and len(res.daily) > 1:
        return [(f"{d:%d/%m}", v[table_title].get(col, 0)) for d, v in res.daily.items()], group
    return [(k, n) for k, n in counts_of(main[2])[table_title].items() if k != "Total"], "Currency"


def custom_content(res: RunResult, opts: dict, bank=None) -> dict:
    """Everything a custom report shows, from the run and the options chosen in the app:
    opts = {"type", "sections": [keys], "views": {key: {"as", "group", "currency"}}, "orientation", "density",
            "include": {"kpis", "currency", "timestamp": True/False}}"""
    bank = get_bank(bank)
    c = bank.COLUMNS
    period = res.period
    paid_on, amount, ref = c["payment_date"], c["amount"], c["reference"]
    main = next((b for b in res.blocks if b[0] in ("Summary-Monthly", "Summary-Daily")), res.blocks[0])
    weeks = [b for b in res.blocks if b[0] == "Summary-Weekly"] if period.report_type == "month" else []
    inc = opts["include"]
    out = {"title": f"{bank.FULL_NAME} Reconciliation Report", "period": period.name[0].upper() + period.name[1:],
           "type": opts["type"], "orientation": opts["orientation"], "density": opts["density"],
           "stamp": f"Reconciled on {res.ran_at:%d %b %Y · %H:%M}" if inc.get("timestamp") else "",
           "kpis": [], "currency_lines": [], "exceptions_line": "", "sections": [],
           "file_stem": f"{bank.CODE}_{opts['type'].lower().replace(' ', '_')}_{period.label}"}
    chosen = [k for k in CUSTOM_SECTIONS if k in opts["sections"]]

    # ---- the overview: key figures and one line per currency
    if "overview" in chosen:
        fig = report_figures(res.blocks)                  # the report's own figures (summary tables)
        if inc.get("kpis"):
            out["kpis"] = [("Processed transaction", f"{fig['processed']:,}"),
                           ("Received notification", f"{fig['received']:,}"),
                           ("Received notification rate", f"{fig['rate']:.2%}" if fig["rate"] is not None else "N/A")]
        if inc.get("currency"):
            out["currency_lines"] = [f"{cur} · Processed transaction {p:,} · Received notification {r:,}"
                                     for cur, p, r in fig["currencies"]]

    # ---- the other sections: a summary table or a chart, and the list of invoices when there is one
    details = res.sheets.get("Payment details", pd.DataFrame())
    base = {paid_on: paid_on, "Currency": "Currency", "invoice_number": "Invoice", amount: "Amount"}
    for key in chosen:
        if key == "overview":
            continue
        title, table_title, views = CUSTOM_SECTIONS[key]
        view = {"as": views[0], "group": chart_groups(period)[0], "currency": "All", **opts["views"].get(key, {})}
        if view["as"] not in views:
            view["as"] = views[0]
        sec = {"key": key, "title": title, "as": view["as"], "table": None, "series": [], "group": "",
               "currency": view["currency"], "list": None, "list_title": ""}
        if key == "merchants":                            # the Payment per merchant table of the main block
            m = res.merchants.get(main[1])
            if m is not None and len(m):
                if view["as"] == TABLE:
                    sec["table"] = pd.DataFrame({"Merchant": m["Merchant name"].values, "Currency": m["Currency"].values,
                                                 "Transactions": m["Number of Transactions"].astype(int).values,
                                                 "Total amount": m["Total Amount"].values,
                                                 "% of received": [f"{v:.2%}" for v in m["%"]]})
                else:                                     # the 10 merchants with the most transactions
                    part = m[m["Merchant name"].ne("All")]
                    if view["currency"] != "All":
                        part = part[part["Currency"].eq(view["currency"])]
                    top = (part.groupby("Merchant name")["Number of Transactions"].sum()
                           .sort_values(ascending=False).head(10))
                    sec["series"], sec["group"] = [(str(k), int(v)) for k, v in top.items()], "Merchant"
            out["sections"].append(sec)
            continue
        if table_title and view["as"] == TABLE:
            sec["table"] = _summary_frame(main[2][table_title])
        elif table_title:
            sec["series"], sec["group"] = _series(res, main, weeks, table_title, view)
        # the invoice list goes with the "Table" view only: a section shown as a chart shows the chart alone
        if key == "pending" and view["as"] == TABLE:
            sec["list"] = _list_frame(res.sheets.get("Pending notification", pd.DataFrame()),
                                      {**base, "payment_status": "Status", "name": "Merchant"})
            sec["list_title"] = "Pending invoices"
        elif key == "other_channel" and view["as"] == TABLE:
            sec["list"] = _list_frame(res.sheets.get("Duplicate payment other channel", pd.DataFrame()),
                                      {**base, "payment_provider": "Provider", "payment_status": "Status"})
            sec["list_title"] = "Invoices paid with another channel"
        elif key == "duplicates":
            if "Debited more than once" in res.sheets:                       # Ecobank's own sheet
                sec["list"] = _list_frame(res.sheets["Debited more than once"],
                                          {**base, "Cba Ref No": "Bank reference", "Debit": "Debit"})
            elif "times_on_statement" in details:
                sec["list"] = _list_frame(details[details["times_on_statement"] > 1],
                                          {**base, ref: "Bank reference", "times_on_statement": "Times on statement"})
            sec["list_title"] = "Invoices credited more than once"
        out["sections"].append(sec)
    return out


def layout_rows(sections: list) -> list:
    """How the sections are placed (preview and PDF): charts that follow each other go side by side, two per row,
    each in its own card; the other sections take the whole width.
    -> [("charts", [section, section]) or ("charts", [section]) or ("section", section)]"""
    rows, run = [], []
    for sec in sections + [None]:
        if sec is not None and sec["series"]:
            run.append(sec)
            continue
        rows += [("charts", run[i:i + 2]) for i in range(0, len(run), 2)]
        run = []
        if sec is not None:
            rows.append(("section", sec))
    return rows


def series_label(sec: dict) -> str:
    """'Transactions by week' (+ ' · RWF' when one currency is chosen)."""
    return (f"Transactions by {sec['group'].lower()}"
            + (f" · {sec['currency']}" if sec["currency"] != "All" and sec["group"] != "Currency" else ""))


def cell_text(name: str, value) -> str:
    """One value of a custom report table as text (PDF and the app's preview): amounts, %, dates, counts."""
    if value is None or (not isinstance(value, str) and pd.isna(value)):
        return ""
    if name == "% of processed":
        return f"{value:.2%}"
    if isinstance(value, pd.Timestamp):
        return f"{value:%d/%m/%Y %H:%M}" if (value.hour or value.minute) else f"{value:%d/%m/%Y}"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return f"{value:,.0f}"
    return str(value)


# ---- the custom report as Excel: the same look as the PDF (IremboPay teal band, KPI cards, styled tables) and the
#      charts as real Excel charts, drawn from the numbers written on the section's sheet
XL_TEAL, XL_GREEN, XL_WARN = "123239", "99E378", "D4A72C"
XL_INK, XL_MUTED, XL_RULE, XL_HAIR = "1D2B2F", "5B6B70", "9AA8AD", "E1E8EA"
XL_TITLE_FILL, XL_TOTAL_FILL, XL_WARN_FILL = "E8F0F2", "F2F6F7", "FFF3D6"
NO_AMOUNT_XL = "—"            # the Total row has no amount: currencies are never added together


def _xl_fill(color: str) -> PatternFill:
    return PatternFill("solid", start_color=color, end_color=color)


def _xl_band(ws, title: str, sub: str, last_col: int, fs: float, back: str = ""):
    """The teal title band across the top of a sheet (+ a link back to the Overview)."""
    fill = _xl_fill(XL_TEAL)
    for r in (1, 2, 3):
        for c in range(1, last_col + 1):
            ws.cell(r, c).fill = fill
    ws.row_dimensions[1].height, ws.row_dimensions[2].height, ws.row_dimensions[3].height = 8, 26, 20
    ws.cell(2, 2, title).font = Font(size=fs + 7, bold=True, color="FFFFFF")
    ws.cell(3, 2, sub).font = Font(size=fs, color=XL_GREEN)
    if back:
        cell = ws.cell(2, last_col, "← Overview")
        cell.hyperlink, cell.font = f"#'{back}'!A1", Font(size=fs, color="FFFFFF", underline="single")
        cell.alignment = Alignment(horizontal="right")
    ws.column_dimensions["A"].width = 2


def _xl_page(ws, content: dict):
    """Print setup: the orientation chosen in the app, one page wide, no grid lines on screen."""
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = content["orientation"].lower()
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = f"{content['title']} · {content['period']} · Page &P of &N"


def _xl_table(ws, top: int, left: int, df: pd.DataFrame, fs: float, title: str = "", total_row: bool = False,
              warning: bool = False, filters: bool = False) -> int:
    """A table like the PDF's: optional title row, tinted header with a teal rule, hairlines between rows, numbers
    on the right, bold tinted Total row (its amount is '—': currencies are never added together).
    Returns the first free row after it."""
    line = XL_WARN if warning else XL_TEAL
    hair, rule = Side(style="thin", color=XL_HAIR), Side(style="medium", color=line)
    cols = list(df.columns)
    numeric = {c for c in cols if pd.api.types.is_numeric_dtype(df[c]) or c == "% of processed"
               or (len(df) and df[c].astype(str).str.endswith("%").all())}
    if title:
        cell = ws.cell(top, left, title)
        cell.font = Font(size=fs + 1, bold=True, color=XL_INK)
        cell.border = Border(left=Side(style="thick", color=line))
        top += 1
    for j, name in enumerate(cols):
        cell = ws.cell(top, left + j, name)
        cell.font = Font(size=fs - 0.5, bold=True, color=XL_MUTED)
        cell.fill = _xl_fill(XL_WARN_FILL if warning else XL_TITLE_FILL)
        cell.border = Border(bottom=rule, top=hair)
        cell.alignment = Alignment(horizontal="right" if name in numeric else "left", vertical="center",
                                   wrap_text=True)
        width = max([len(str(name))] + [len(cell_text(name, v)) for v in df[name].head(300)])
        letter = get_column_letter(left + j)
        ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width or 0, min(max(width + 3, 11), 45))
    ws.row_dimensions[top].height = 20
    for i, rec in enumerate(df.itertuples(index=False), 1):
        is_total = total_row and str(rec[0]) in ("Total", "All")
        for j, (name, value) in enumerate(zip(cols, rec)):
            empty = value is None or (not isinstance(value, str) and pd.isna(value))
            if empty and is_total and name == "Total amount":
                value, empty = NO_AMOUNT_XL, False
            cell = ws.cell(top + i, left + j, None if empty else (value.to_pydatetime() if isinstance(value, pd.Timestamp)
                                                                else value))
            cell.font = Font(size=fs, bold=is_total, color="7A4E00" if warning else XL_INK)
            cell.border = Border(bottom=hair, top=rule if is_total else None)
            if is_total:
                cell.fill = _xl_fill(XL_TOTAL_FILL)
            cell.alignment = Alignment(horizontal="right" if name in numeric else "left", vertical="center")
            if name == "% of processed":
                cell.number_format = "0.00%"
            elif isinstance(value, pd.Timestamp):
                cell.number_format = "dd/mm/yyyy hh:mm" if (value.hour or value.minute) else "dd/mm/yyyy"
            elif isinstance(value, (int, float)) and not isinstance(value, bool):
                cell.number_format = "#,##0"
    last = top + len(df)
    if filters and len(df):
        ws.auto_filter.ref = f"{get_column_letter(left)}{top}:{get_column_letter(left + len(cols) - 1)}{last}"
    return last + 2



def _xl_chart(sec: dict, ws, head: int, n: int):
    """A one-series Excel bar or line chart of the section's numbers (written on `ws` at rows head..head+n, columns
    B:C), teal like the PDF (amber for SLA), values on the bars when there are few, no legend."""
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.chart.label import DataLabelList
    color = XL_WARN if sec["key"] == "sla" else XL_TEAL
    ch = BarChart() if sec["as"] == BAR else LineChart()
    if sec["as"] == BAR:
        ch.type, ch.gapWidth = "col", 60
    ch.add_data(Reference(ws, min_col=3, min_row=head, max_row=head + n), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, min_row=head + 1, max_row=head + n))
    s = ch.series[0]
    s.graphicalProperties.line.solidFill = color
    if sec["as"] == BAR:
        s.graphicalProperties.solidFill = color
    else:
        s.graphicalProperties.line.width = 28575                          # 2.25 pt
        s.marker.symbol, s.marker.size = "circle", 6
        s.marker.graphicalProperties.solidFill = color
        s.marker.graphicalProperties.line.solidFill = "FFFFFF"
        s.smooth = False
    if n <= 12:
        s.dLbls = DataLabelList()
        s.dLbls.showVal = True
        for flag in ("showSerName", "showCatName", "showLegendKey", "showPercent"):
            setattr(s.dLbls, flag, False)
        if sec["as"] == LINE:
            s.dLbls.position = "t"
    ch.title = f"{sec['title']} · {series_label(sec)}"
    ch.legend = None
    ch.y_axis.title = "Transactions"
    ch.y_axis.number_format = "#,##0"
    ch.x_axis.delete = ch.y_axis.delete = False                          # Excel hides the axes otherwise
    ch.y_axis.majorGridlines.spPr = None
    ch.width, ch.height = 16, 8
    return ch


def custom_excel(content: dict) -> bytes:
    """The custom report as an Excel file, laid out like the PDF:
    Overview - title band, key figures as cards, the line per currency, a linked list of the sections and every
               chosen chart (two per row, real Excel charts);
    then one sheet per section - its table, or the numbers behind its chart with the chart beside them, and its
               invoice list (with filters). Orientation and density chosen in the app are kept for printing."""
    from openpyxl import Workbook
    fs = 9 if content["density"] == "Compact" else 10
    sheet_of = {sec["key"]: sec["title"][:31] for sec in content["sections"]}
    wb = Workbook()
    ov = wb.active
    ov.title = "Overview"
    _xl_page(ov, content)
    for c in range(2, 18):
        ov.column_dimensions[get_column_letter(c)].width = 11
    _xl_band(ov, content["title"], f"{content['period']}  ·  {content['type']}"
             + (f"  ·  {content['stamp']}" if content["stamp"] else ""), 17, fs)
    row = 5

    # ---- key figures: one card each (label over a large value, green bar on the left), like the PDF
    if content["kpis"]:
        side = Side(style="thin", color=XL_RULE)
        for i, (label, value) in enumerate(content["kpis"]):
            c0 = 2 + i * 4
            for r in (row, row + 1):
                ov.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + 2)
                for c in range(c0, c0 + 3):
                    ov.cell(r, c).border = Border(top=side if r == row else None, bottom=side if r > row else None,
                                                  left=Side(style="thick", color=XL_GREEN) if c == c0 else None,
                                                  right=side if c == c0 + 2 else None)
            ov.cell(row, c0, label.upper()).font = Font(size=fs - 1, bold=True, color=XL_MUTED)
            ov.cell(row + 1, c0, value).font = Font(size=fs + 8, bold=True, color=XL_TEAL)
            for r in (row, row + 1):
                ov.cell(r, c0).alignment = Alignment(indent=1, vertical="center")
        ov.row_dimensions[row].height, ov.row_dimensions[row + 1].height = 18, 30
        row += 3
    lines = content["currency_lines"] + ([content["exceptions_line"]] if content["exceptions_line"] else [])
    for line in lines:
        for c in range(2, 14):
            ov.cell(row, c).fill = _xl_fill(XL_TOTAL_FILL)
        ov.cell(row, 2, line).font = Font(size=fs, color=XL_INK)
        row += 1
    row += 1 if lines else 0

    # ---- the sections of this report, each a link to its sheet
    if content["sections"]:
        ov.cell(row, 2, "Sections").font = Font(size=fs + 2, bold=True, color=XL_INK)
        row += 1
        for sec in content["sections"]:
            shown = sec["as"] if sec["as"] != TABLE else "Table"
            cell = ov.cell(row, 2, sec["title"])
            cell.hyperlink, cell.font = f"#'{sheet_of[sec['key']]}'!A1", Font(size=fs, color="0563C1",
                                                                             underline="single")
            ov.cell(row, 5, shown).font = Font(size=fs, color=XL_MUTED)
            if sec["list"] is not None:
                ov.cell(row, 7, f"{len(sec['list']):,} invoices").font = Font(size=fs, color=XL_MUTED)
            row += 1
        row += 1

    # ---- one sheet per section; the charts are also placed on the Overview, two per row
    charts = []
    for sec in content["sections"]:
        s = wb.create_sheet(sheet_of[sec["key"]])
        _xl_page(s, content)
        width = max(8, len(sec["list"].columns) + 1 if sec["list"] is not None else 0,
                    len(sec["table"].columns) + 1 if sec["table"] is not None else 0)
        _xl_band(s, sec["title"], series_label(sec) if sec["series"] else content["period"], width + 1, fs,
                 back="Overview")
        r = 5
        if sec["table"] is not None:
            r = _xl_table(s, r, 2, sec["table"], fs, title="Summary", total_row=True, warning=sec["key"] == "sla")
        if sec["series"]:
            data = pd.DataFrame(sec["series"], columns=[sec["group"], "Transactions"])
            _xl_table(s, r, 2, data, fs, title="Numbers behind the chart")
            head = r + 1
            s.add_chart(_xl_chart(sec, s, head, len(data)), f"E{r}")
            charts.append(_xl_chart(sec, s, head, len(data)))
            r = max(r + len(data) + 3, r + 18)
        if sec["list"] is not None:
            title = f"{sec['list_title']} ({len(sec['list']):,})"
            if len(sec["list"]):
                _xl_table(s, r, 2, sec["list"], fs, title=title, filters=True)
                s.freeze_panes = s.cell(r + 2, 1)
            else:
                s.cell(r, 2, title).font = Font(size=fs + 1, bold=True, color=XL_INK)
                s.cell(r + 1, 2, "None.").font = Font(size=fs, color=XL_MUTED)
    if charts:
        ov.cell(row, 2, "Charts").font = Font(size=fs + 2, bold=True, color=XL_INK)
        row += 1
        for i, ch in enumerate(charts):
            ov.add_chart(ch, f"{'B' if i % 2 == 0 else 'J'}{row + (i // 2) * 17}")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _nice_step(top: float) -> float:
    """A round step for a chart axis with about 4 lines (1, 2, 5, 10, 20, ...)."""
    raw = max(top / 4, 1)
    mag = 10 ** int(len(str(int(raw))) - 1)
    return next(m * mag for m in (1, 2, 5, 10) if m * mag >= raw)


def _pdf_chart(kind: str, points: list, width: float, height: float):
    """A one-series bar or line chart (IremboPay teal, light grid, no legend: the section title names it)."""
    from reportlab.graphics.charts.barcharts import VerticalBarChart
    from reportlab.graphics.charts.linecharts import HorizontalLineChart
    from reportlab.graphics.shapes import Drawing
    from reportlab.graphics.widgets.markers import makeMarker
    from reportlab.lib import colors

    teal, ink, grid = colors.HexColor(CHART_TEAL), colors.HexColor("#5B6B70"), colors.HexColor("#E1E7E9")
    labels, values = [p[0] for p in points] or ["–"], [p[1] for p in points] or [0]
    many = len(labels) > 12
    width = min(width, 120 + 70 * len(labels))           # few bars: a narrower chart, not very wide bars
    d = Drawing(width, height)
    d.hAlign = "LEFT"
    ch = VerticalBarChart() if kind == BAR else HorizontalLineChart()
    ch.x, ch.y = 36, 30 if many else 20
    ch.width, ch.height = width - 46, height - ch.y - 8
    ch.data = [values]
    step = _nice_step(max(values))
    ch.valueAxis.valueMin, ch.valueAxis.valueStep = 0, step
    ch.valueAxis.valueMax = step * (int(max(values) / step) + 1)
    ch.valueAxis.visibleGrid, ch.valueAxis.gridStrokeColor, ch.valueAxis.gridStrokeWidth = True, grid, 0.5
    ch.valueAxis.strokeColor, ch.valueAxis.visibleTicks = grid, False
    ch.valueAxis.labels.fontSize, ch.valueAxis.labels.fillColor = 7, ink
    ch.valueAxis.labels.fontName = ch.categoryAxis.labels.fontName = "Helvetica"
    ch.valueAxis.labelTextFormat = lambda v: f"{int(v):,}"
    step_l = 1 if ch.width / len(labels) >= 14 else 2                 # narrow chart: every second label
    ch.categoryAxis.categoryNames = [l if i % step_l == 0 else "" for i, l in enumerate(labels)]
    ch.categoryAxis.strokeColor, ch.categoryAxis.visibleTicks = grid, False
    ch.categoryAxis.labels.fontSize, ch.categoryAxis.labels.fillColor = 7, ink
    if many:
        ch.categoryAxis.labels.angle, ch.categoryAxis.labels.boxAnchor, ch.categoryAxis.labels.dy = 90, "e", -2
    if kind == BAR:
        ch.bars[0].fillColor, ch.bars[0].strokeColor = teal, None
        ch.groupSpacing = 2 if many else 8
    else:
        ch.lines[0].strokeColor, ch.lines[0].strokeWidth = teal, 2
        ch.lines[0].symbol = makeMarker("FilledCircle", size=5, fillColor=teal, strokeColor=colors.white,
                                        strokeWidth=1)
    d.add(ch)
    return d


def custom_pdf(content: dict) -> bytes:
    """The custom report as a PDF: title band, overview, then each section (table, chart and invoice list)."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import KeepTogether, LongTable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape

    teal, green = colors.HexColor("#123239"), colors.HexColor("#99E378")
    title_fill, grid, muted = colors.HexColor("#D0DEE3"), colors.HexColor("#9AA8AD"), colors.HexColor("#5B6B70")
    size = landscape(A4) if content["orientation"] == "Landscape" else A4
    page_w = size[0]
    margin = 14 * mm
    usable = page_w - 2 * margin
    compact = content["density"] == "Compact"
    fs, pad, gap = (7, 1.5, 4 * mm) if compact else (8.5, 3, 7 * mm)
    h1 = ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=16, textColor=green, leading=20)
    sub = ParagraphStyle("sub", fontName="Helvetica", fontSize=9.5, textColor=colors.HexColor("#D6E4E6"), leading=12)
    h2 = ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=12 if compact else 13, textColor=teal, leading=16,
                        spaceAfter=3)
    h3 = ParagraphStyle("h3", fontName="Helvetica-Bold", fontSize=fs + 1, textColor=teal, leading=fs + 4,
                        spaceBefore=4, spaceAfter=2)
    body = ParagraphStyle("body", fontName="Helvetica", fontSize=fs + 0.5, leading=fs + 4)
    small = ParagraphStyle("small", fontName="Helvetica", fontSize=fs, textColor=muted, leading=fs + 3)
    cell = ParagraphStyle("cell", fontName="Helvetica", fontSize=fs, leading=fs + 2)

    def table(df: pd.DataFrame, width: float, total_row=False):
        """A table: bold header, light grid, numbers on the right; long text wraps."""
        cols = list(df.columns)
        rows = [cols] + [[Paragraph(escape(cell_text(c, v)), cell) if c in ("Merchant", "Comment", "Decision")
                          else cell_text(c, v) for c, v in zip(cols, rec)] for rec in df.itertuples(index=False)]
        weights = [max([len(str(c))] + [len(cell_text(c, v)) for v in df[c].head(100)]) for c in cols]
        weights = [min(max(w, 6), 40) for w in weights]
        widths = [width * w / sum(weights) for w in weights]
        numeric = [j for j, c in enumerate(cols) if pd.api.types.is_numeric_dtype(df[c]) or c == "% of processed"]
        t = LongTable(rows, colWidths=widths, repeatRows=1, hAlign="LEFT")
        style = [("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("BACKGROUND", (0, 0), (-1, 0), title_fill),
                 ("FONTSIZE", (0, 0), (-1, -1), fs), ("GRID", (0, 0), (-1, -1), 0.4, grid),
                 ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                 ("TOPPADDING", (0, 0), (-1, -1), pad), ("BOTTOMPADDING", (0, 0), (-1, -1), pad)]
        style += [("ALIGN", (j, 1), (j, -1), "RIGHT") for j in numeric]
        if total_row:
            style.append(("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"))
        t.setStyle(TableStyle(style))
        return t

    # ---- title band + overview
    band = Table([[Paragraph(escape(content["title"]), h1)],
                  [Paragraph(escape(" · ".join(x for x in (content["period"], content["type"], content["stamp"]) if x)),
                             sub)]], colWidths=[usable])
    band.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), teal), ("LEFTPADDING", (0, 0), (-1, -1), 10),
                              ("TOPPADDING", (0, 0), (0, 0), 8), ("BOTTOMPADDING", (0, -1), (-1, -1), 8),
                              ("LINEBEFORE", (0, 0), (0, -1), 5, green)]))
    story = [band]
    if content["kpis"]:
        n = len(content["kpis"])
        k = Table([[v for _, v in content["kpis"]], [label for label, _ in content["kpis"]]], colWidths=[usable / n] * n)
        k.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                               ("FONTSIZE", (0, 0), (-1, 0), 13 if compact else 16), ("TEXTCOLOR", (0, 0), (-1, 0), teal),
                               ("FONTSIZE", (0, 1), (-1, 1), fs), ("TEXTCOLOR", (0, 1), (-1, 1), muted),
                               ("ALIGN", (0, 0), (-1, -1), "CENTER"), ("BOX", (0, 0), (-1, -1), 0.6, grid),
                               ("LINEBEFORE", (1, 0), (-1, -1), 0.6, grid),
                               ("TOPPADDING", (0, 0), (-1, 0), 4 if compact else 7),
                               ("BOTTOMPADDING", (0, 1), (-1, 1), 4 if compact else 7)]))
        story += [Spacer(1, gap), k]
    lines = content["currency_lines"] + ([content["exceptions_line"]] if content["exceptions_line"] else [])
    if lines:
        story += [Spacer(1, gap / 2)] + [Paragraph(escape(line), body) for line in lines]

    def invoice_list(sec, named=False):
        """The invoices of a section, under its table, or under its chart card (named: the section is said)."""
        if sec["list"] is None:
            return []
        title = (f"{sec['title']} · " if named else "") + f"{sec['list_title']} ({len(sec['list']):,})"
        return [Paragraph(escape(title), h3),
                table(sec["list"], usable) if len(sec["list"]) else Paragraph("None.", small)]

    def chart_card(sec, width):
        """One chart in its own card: title, what is counted, the chart, the values (for print)."""
        inner = width - 16
        card = Table([[[Paragraph(escape(sec["title"]), h3), Paragraph(escape(series_label(sec)), small),
                        _pdf_chart(sec["as"], sec["series"], inner, (42 if compact else 55) * mm),
                        Paragraph(escape(" · ".join(f"{name} {v:,}" for name, v in sec["series"])), small)]]],
                     colWidths=[width], hAlign="LEFT")
        card.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.6, grid), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                  ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                                  ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
        return card

    # ---- the sections: charts side by side (two per row, each in a card), tables and lists full width
    space = 5 * mm                                      # between two chart cards
    for kind, item in layout_rows(content["sections"]):
        if kind == "section":
            sec = item
            part = [Spacer(1, gap), Paragraph(escape(sec["title"]), h2)]
            if sec["table"] is not None:
                part.append(table(sec["table"], min(usable, 150 * mm), total_row=True))
            story.append(KeepTogether(part))
            story += invoice_list(sec)
            continue
        half = (usable - space) / 2
        cards = [chart_card(sec, half if len(item) == 2 else usable) for sec in item]
        row = Table([cards], colWidths=[half + space, half] if len(cards) == 2 else [usable], hAlign="LEFT")
        row.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                                 ("RIGHTPADDING", (0, 0), (0, -1), space if len(cards) == 2 else 0),
                                 ("RIGHTPADDING", (1, 0), (-1, -1), 0),
                                 ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0)]))
        story += [Spacer(1, gap), row]
        for sec in item:
            story += invoice_list(sec, named=True)

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(muted)
        canvas.drawRightString(page_w - margin, 8 * mm, f"Page {doc.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=size, leftMargin=margin, rightMargin=margin, topMargin=margin,
                            bottomMargin=margin, title=f"{content['title']} {content['period']}")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()
