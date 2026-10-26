"""BPR Bank Rwanda: made-up statements in the BPR export layout (no real account or invoice data). The rules decided by
the user on 16/10/2026: DOC ID = invoice, VALUE_DATE = date, currency from IremboPay, BPR_RW only, SLA in whole minutes
(1 by default), found-but-doubtful credits to review, files checked against their balances and each other."""
import io

import pandas as pd
import pytest
from openpyxl import Workbook

import engine as e
from banks import bpr
from services import app_settings

HEAD = ["BOOKING_DATE", "BANK_REFERENCE", "DESCRIPTION", "DOC ID", "VALUE_DATE", "DEBIT", "CREDIT", "CLOSING_BALANCE"]


def statement(rows, opening=1000, header=HEAD, sheets=None):
    """rows: (booking, ref, description, doc_id, value_date, debit, credit); the balances are computed."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Export Worksheet"
    ws.append(header)
    ws.append([rows[0][0] if rows else "2026-08-01", "", "Opening balance", "", "", 0, 0, opening])
    bal = opening
    for b, ref, desc, doc, val, debit, credit in rows:
        bal += credit + debit
        ws.append([b, ref, desc, doc, val, debit, credit, bal])
    for name, content in (sheets or {}).items():
        s = wb.create_sheet(name)
        for r in content:
            s.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def payments(rows):
    cols = ["invoice_number", "transaction_id", "amount", "currency", "payment_status", "description", "expiry_at",
            "created_at", "settled_at", "paid_at", "name", "transaction_reference", "provider_reference",
            "payment_channel", "payment_provider", "transaction_status"]
    out = [",".join(cols)]
    for inv, amount, cur, paid, provider, pref in rows:
        out.append(",".join([inv, "T" + inv, str(amount), cur, "SETTLED", "", "", paid, "", paid, "RSSB", "",
                             pref, "BPR_RW_AGENT", provider, "SUCCESSFUL"]))
    return "\n".join(out).encode()


def read(content, name="aug.xlsx"):
    return e.read_bank_file(name, content, "BPR")


AUG = e.make_period("month", "2026-08-01")
ROWS = [  # booking, ref, description, DOC ID, VALUE_DATE, debit, credit
    ("2026-08-03", "FT1", "SZAAA1", "886000000001", "2026-08-03 10:00:00", 0, 5000),     # on time (same minute)
    ("2026-08-03", "FT2", "20140000002", "886000000002", "2026-08-03 11:00:00", 0, 2000),  # 2 minutes late
    ("2026-08-03", "FT3", "SZAAA3", "", "2026-08-03 12:00:00", 0, 700),                 # no DOC ID
    ("2026-08-03", "FT4", "SZAAA4", "886000000004", "2026-08-03 13:00:00", 0, 300),      # other provider
    ("2026-08-03", "FT5", "SZAAA5", "886000000005", "2026-08-03 14:00:00", 0, 1000),     # credited twice: the
    ("2026-08-03", "FT6", "SZAAA6", "886000000005", "2026-08-03 14:05:00", 0, 1000),     # IremboPay one is FT6
    ("2026-08-03", "FT7", "886000000007", "886000000007", "2026-08-03 15:00:00", 0, 150),  # USD (from IremboPay)
    ("2026-08-04", "RWF1", "", "", "2026-08-04 03:00:00", -9000, 0),                    # daily transfer out
]
PAYS = [("886000000001", 5000, "RWF", "2026-08-03 10:00:40", "BPR_RW", "SZAAA1"),
        ("886000000002", 2000, "RWF", "2026-08-03 11:02:10", "BPR_RW", "20140000002"),
        ("886000000004", 300, "RWF", "2026-08-03 13:00:20", "MTN", ""),
        ("886000000005", 1000, "RWF", "2026-08-03 14:05:30", "BPR_RW", "SZAAA6"),
        ("886000000007", 150, "USD", "2026-08-03 15:01:00", "BPR_RW", "886000000007")]


def run(tmp_path, rows=ROWS, pays=PAYS, require_review=True, sla=1):
    files = [read(statement(rows)), e.read_payment_file("p.csv", payments(pays), "BPR")]
    return e.run(tmp_path / "base", AUG, files, sla_days=sla, require_review=require_review, bank="BPR",
                 report_dir=tmp_path / "r")


# ---- reading --------------------------------------------------------------------------------------------------------
def test_statement_is_read_whatever_the_doc_id_title():
    for title in ("DOC ID", "Doc ID", "doc id", "DOCID", ""):
        f = read(statement(ROWS, header=HEAD[:3] + [title] + HEAD[4:]))
        assert f.rows == 7 and f.currency == "" and f.bank == "BPR"           # the 7 credits; the debit is not kept
        assert bpr.clean_statement(f.data)["invoice_number"].iloc[0] == "886000000001"


def test_the_data_sheet_is_found_and_other_sheets_ignored():
    content = statement(ROWS, sheets={"Sheet1": [["20145323788", "886244898333"]], "SQL": [["select ..."]]})
    assert read(content).rows == 7


def test_a_file_that_does_not_add_up_is_refused():
    wb = Workbook()
    ws = wb.active
    ws.append(HEAD)
    ws.append(["2026-08-03", "", "Opening balance", "", "", 0, 0, 1000])
    ws.append(["2026-08-03", "FT1", "SZ1", "886000000001", "2026-08-03 10:00:00", 0, 500, 1500])
    ws.append(["2026-08-03", "FT2", "SZ2", "886000000002", "2026-08-03 11:00:00", 0, 500, 1900])   # should be 2000
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(e.ReconError, match="do not add up"):
        read(buf.getvalue())


def test_without_value_date_the_booking_date_is_used_and_the_row_is_flagged():
    head = [h for h in HEAD if h != "VALUE_DATE"]
    rows = [r[:4] + r[5:] for r in ROWS[:1]]
    wb = Workbook()
    ws = wb.active
    ws.append(head)
    ws.append(["2026-08-03", "", "Opening balance", "", 0, 0, 1000])
    ws.append(list(rows[0]) + [6000])
    buf = io.BytesIO()
    wb.save(buf)
    st = bpr.clean_statement(read(buf.getvalue()).data)
    assert st["VALUE_DATE"].iloc[0] == pd.Timestamp("2026-08-03") and st[bpr.NO_VALUE_DATE].iloc[0] == "yes"


def test_a_credit_also_in_a_file_with_value_date_is_kept_once():
    dated = read(statement([("2026-10-12", "FT9", "SZ9", "886000000009", "2026-10-12 09:41:00", 0, 500)]), "oct 12.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.append([h for h in HEAD if h != "VALUE_DATE"])
    ws.append(["2026-09-21", "", "Opening balance", "", 0, 0, 1000])
    ws.append(["2026-10-12", "FT9", "SZ9", "886000000009", 0, 500, 1500])
    ws.append(["2026-10-12", "FT8", "SZ8", "886000000008", 0, 700, 2200])
    buf = io.BytesIO()
    wb.save(buf)
    undated = read(buf.getvalue(), "sept 22 to oct 13.xlsx")
    st = bpr.clean_statement(pd.concat([dated.data, undated.data], ignore_index=True))
    assert sorted(st["BANK_REFERENCE"]) == ["FT8", "FT9"]
    assert st.loc[st["BANK_REFERENCE"].eq("FT9"), bpr.NO_VALUE_DATE].iloc[0] == ""        # the dated copy is kept


def test_a_credit_cancelled_by_a_debit_is_flagged():
    rows = [("2026-10-05", "TT1", "886000000009", "", "2026-10-05 15:00:00", -6000, 0),
            ("2026-10-05", "TT1", "886000000009", "", "2026-10-05 15:00:00", 0, 6000)]
    data = read(statement(rows)).data
    assert data[bpr.CANCELLED].iloc[0] == "TT1"


# ---- the checks between files (step 3) ------------------------------------------------------------------------------
def test_files_must_follow_each_other():
    day = lambda d, credit, opening: read(statement([(d, f"FT{d}", "SZ", "886000000001", f"{d} 10:00:00", 0, credit)],
                                                    opening=opening), f"{d}.xlsx")
    a, b, c = day("2026-08-01", 500, 1000), day("2026-08-03", 200, 1500), day("2026-08-04", 100, 1700)
    covered, lines, blocking = bpr.check_files([a, b, c])
    assert blocking == [] and pd.Timestamp("2026-08-02") in covered          # 2 Aug: no payment, the balance follows
    shifted = day("2026-08-03", 200, 1520)                                   # its balance column is shifted by +20
    covered, lines, blocking = bpr.check_files([a, shifted, c])
    assert blocking == [] and any("shifted by +20" in t for ok, t in lines)
    gap = day("2026-08-04", 100, 1900)                                       # 200 came in with no row
    covered, lines, blocking = bpr.check_files([a, b, gap])
    assert blocking == [] and any("a payment is missing" in t for ok, t in lines if not ok)   # shown, not blocking
    out = day("2026-08-04", 100, 1600)                                       # 100 went out with no row
    covered, lines, blocking = bpr.check_files([a, b, out])
    assert blocking == [] and any("a debit is missing" in t for ok, t in lines if not ok)


def test_a_file_repeating_the_last_row_of_the_file_before_follows_it():
    # 15 Oct starts with the 14 Oct debit again (and no Opening balance row): no payment is missing
    d14 = read(statement([("2026-10-14", "FT1", "SZ1", "886000000001", "2026-10-14 10:00:00", 0, 500),
                          ("2026-10-14", "RWF14", "", "", "2026-10-14 03:44:00", -1200, 0)], opening=1000), "oct 14.xlsx")
    d15 = read(statement([("2026-10-14", "RWF14", "", "", "2026-10-14 03:44:00", -1200, 0),
                          ("2026-10-15", "FT2", "SZ2", "886000000002", "2026-10-15 10:00:00", 0, 300)],
                         opening=1500), "OCT 15 RWF.xlsx")
    covered, lines, blocking = bpr.check_files([d14, d15])
    assert all(ok for ok, t in lines) and any("already in oct 14.xlsx" in t for ok, t in lines)


def test_a_file_inside_another_is_left_out_of_the_check():
    rows = [("2026-08-06", "TT1", "", "886000000001", "2026-08-06 12:00:00", 0, 750),
            ("2026-08-07", "RWF1", "", "", "2026-08-07 03:00:00", -750, 0)]
    whole, part = read(statement(rows, opening=4), "1 to 10 USD.xlsx"), read(statement(rows[:1], opening=4), "6 USD.xlsx")
    covered, lines, blocking = bpr.check_files([whole, part])
    assert blocking == [] and any("is already inside" in t for ok, t in lines)


def test_step3_recommends_the_day_after_and_lists_payments_on_no_statement(tmp_path):
    files = [read(statement(ROWS)), e.read_payment_file("p.csv", payments(PAYS + [
        ("886000000099", 100, "RWF", "2026-08-10 09:00:00", "BPR_RW", "X")]), "BPR")]
    chk = e.check_period(AUG, files, "BPR", today="2026-10-16")
    assert any("statement of 01/09/2026" in w for w in chk.warnings)
    assert any(t.startswith("Paid in IremboPay through BPR: 1 of") and t.endswith("is not in the BPR statements")
               for ok, t in chk.lines if not ok)
    assert not any("886000000099" in w for w in chk.warnings)          # one message, not repeated in a warning


# ---- the run --------------------------------------------------------------------------------------------------------
def test_what_goes_to_review(tmp_path):
    res = run(tmp_path)
    reasons = dict(zip(res.review["BANK_REFERENCE"], res.review["Reason"]))
    assert "FT3" not in reasons                                        # no invoice number: Pending, not reviewed
    assert reasons["FT4"].startswith("Paid through MTN in IremboPay")
    assert reasons["FT5"] == "Invoice credited again: IremboPay has no record of this credit; the credit counted is FT6"
    assert "FT6" not in reasons and "FT1" not in reasons
    assert (res.not_found, res.to_review) == (1, 2) and res.stopped_for_review


def test_currency_from_irembopay_and_sla_in_whole_minutes(tmp_path):
    res = run(tmp_path, require_review=False)
    main = res.blocks[-1][2]
    usd = main["Overall processed payments"]
    assert int(usd.loc[usd["Currency"] == "USD", "Number of Transactions"].iloc[0]) == 1      # FT7: USD from IremboPay
    late = res.sheets["Payment details"].set_index("BANK_REFERENCE")["sla_status"]
    assert late["FT1"] == "On time"          # paid 10:00:40 -> minute 10:00 = VALUE_DATE
    assert late["FT7"] == "On time"          # 15:01 vs 15:00: 1 minute allowed
    assert late["FT2"] == "SLA violated"     # 11:02 vs 11:00: 2 minutes
    res0 = run(tmp_path / "zero", require_review=False, sla=0)
    assert res0.sheets["Payment details"].set_index("BANK_REFERENCE")["sla_status"]["FT7"] == "SLA violated"


def test_credits_without_invoice_go_straight_to_pending(tmp_path):
    res = run(tmp_path)
    assert "FT3" not in res.review["BANK_REFERENCE"].tolist()          # not in Exception Resolution
    review = res.review.copy()
    review["Reviewed (Yes/No)"] = "Yes"
    e.save_decisions(tmp_path / "base", AUG, review, "BPR")
    res2 = run(tmp_path)
    assert not res2.stopped_for_review and res2.report_file is not None
    pend = res2.sheets["Pending notification"].set_index("BANK_REFERENCE")
    assert pend.loc["FT3", "Currency"] == "RWF"                          # from the file name
    assert pend.loc["FT3", "review_status"] == "Pending notification (no invoice number)"
    from openpyxl import load_workbook
    assert "Credits without DOC ID" not in load_workbook(res2.report_file, read_only=True).sheetnames


def test_processed_is_every_credit_on_the_statement_each_invoice_once(tmp_path):
    # 26/10/2026: processed = every credit (FT5 is the second credit of FT6's invoice: Exception Resolution);
    # received = paid_at filled through BPR_RW; FT3 (no invoice number) is pending; FT4 (MTN) duplicate
    res = run(tmp_path, require_review=False)
    t = e.counts_of(res.blocks[-1][2])
    assert t["Overall processed payments"]["Total"] == 6                       # FT1 FT2 FT3 FT4 FT6 FT7
    assert t["Overall notifications received by Irembo"]["Total"] == 4         # FT1 FT2 FT6 FT7
    assert t["Duplicate Payment"]["Total"] == 1 and t["Pending notifications"]["Total"] == 1     # FT4 / FT3
    assert t["SLA Violated"]["Total"] == 1                                     # FT2
    details = res.sheets["Payment details"].set_index("BANK_REFERENCE")
    assert bool(details.loc["FT5", "is_extra_credit"]) and not bool(details.loc["FT6", "is_extra_credit"])
    assert "FT3" in res.sheets["Pending notification"]["BANK_REFERENCE"].tolist()


def test_step6_and_the_custom_report_show_the_report_figures(tmp_path):
    # 22/10/2026: step 6 and the custom report read the report's summary tables (not the statement rows)
    res = run(tmp_path, require_review=False)
    fig = e.report_figures(res.blocks)
    assert res.bank_rows == 7 and (fig["processed"], fig["received"]) == (6, 4)    # FT5 extra
    assert round(fig["rate"], 4) == round(4 / 6, 4) and fig["currencies"] == [("RWF", 5, 3), ("USD", 1, 1)]
    opts = {"type": "Summary", "orientation": "portrait", "density": "normal", "sections": ["overview"],
            "views": {}, "include": {"kpis": True, "currency": True, "timestamp": False}}
    out = e.custom_content(res, opts, "BPR")
    assert out["kpis"][0] == ("Processed transaction", "6") and out["kpis"][1] == ("Received notification", "4")
    assert out["kpis"][2] == ("Received notification rate", "66.67%")


def test_a_cancelled_credit_without_invoice_is_still_reviewed(tmp_path):
    rows = [("2026-08-05", "TT9", "SZX", "", "2026-08-05 15:00:00", 0, 6000),
            ("2026-08-05", "TT9", "SZX", "", "2026-08-05 15:30:00", -6000, 0)]
    res = run(tmp_path, rows=rows)
    assert res.review["Reason"].tolist() == ["No DOC ID; cancelled by the debit TT9"]


def test_sla_setting_is_in_minutes_with_1_by_default(tmp_path):
    data = app_settings.load(tmp_path / "none.json")
    assert app_settings.sla_days(data, "BPR") == 1 and app_settings.sla_unit("BPR") == "minutes"
    assert app_settings.sla_days(data, "IM") == 0 and app_settings.sla_unit("IM") == "days"
    app_settings.set_sla_days("BPR", 5, "ed@irembo.com", tmp_path / "s.json")
    assert app_settings.sla_days(app_settings.load(tmp_path / "s.json"), "BPR") == 5
    with pytest.raises(app_settings.SettingsError):
        app_settings.set_sla_days("BPR", 61, "ed@irembo.com", tmp_path / "s.json")


def test_credited_again_says_whether_irembopay_recorded_it(tmp_path):
    """Like invoice 886239268688 on 27/08: 4 credits, IremboPay recorded 2 transactions; 1 counts, 3 are checked."""
    rows = [("2026-08-27", f"FT{n}", f"SZ{n}", "886000000077", f"2026-08-27 13:2{n}:00", 0, 10000) for n in range(4)]
    pays = [("886000000077", 10000, "RWF", "2026-08-27 13:28:36", "BPR_RW", "SZ3"),
            ("886000000077", 10000, "RWF", "2026-08-27 13:28:36", "BPR_RW", "SZ1")]
    res = run(tmp_path, rows=rows, pays=pays)
    reasons = dict(zip(res.review["BANK_REFERENCE"], res.review["Reason"]))
    counted = next(iter({"FT0", "FT1", "FT2", "FT3"} - set(reasons)))
    assert len(reasons) == 3 and counted in ("FT1", "FT3")
    other_recorded = ({"FT1", "FT3"} - {counted}).pop()
    assert reasons[other_recorded].startswith("Invoice credited again: IremboPay also recorded this credit")
    assert reasons["FT0"].startswith("Invoice credited again: IremboPay has no record of this credit")


def test_a_doc_id_not_in_irembopay_takes_the_currency_of_the_file_name(tmp_path):
    rows = [("2026-08-03", "FT1", "SZ1", "886000000088", "2026-08-03 10:00:00", 0, 750)]
    for name, cur in (("aug 3 USD.xlsx", "USD"), ("aug 3 rwf.xlsx", "RWF")):
        files = [read(statement(rows), name), e.read_payment_file("p.csv", payments(PAYS), "BPR")]
        res = e.run(tmp_path / name, AUG, files, sla_days=1, bank="BPR", report_dir=tmp_path / "r")
        assert res.review["Currency"].iloc[0] == cur and res.review["Reason"].iloc[0].startswith("DOC ID not in")
        review = res.review.copy()
        review["Reviewed (Yes/No)"], review["Report table"] = "Yes", "Paid with other channel"
        e.save_decisions(tmp_path / name, AUG, review, "BPR")                 # no currency to choose
        res2 = e.run(tmp_path / name, AUG, files, sla_days=1, bank="BPR", report_dir=tmp_path / "r")
        assert not res2.stopped_for_review


def test_decisions_saved_with_a_blank_currency_are_still_found(tmp_path):
    """A decision saved before the currency of a 'DOC ID not in IremboPay' credit came from the file name."""
    rows = [("2026-08-03", "FT1", "SZ1", "886000000088", "2026-08-03 10:00:00", 0, 750)]
    files = [read(statement(rows), "aug 3 rwf.xlsx"), e.read_payment_file("p.csv", payments(PAYS), "BPR")]
    res = e.run(tmp_path, AUG, files, sla_days=1, bank="BPR", report_dir=tmp_path / "r")
    review = res.review.copy()
    review["Currency"] = ""                                  # as it was saved before 16/10/2026 15:30
    review["Reviewed (Yes/No)"], review["Report table"] = "Yes", "Paid with other channel"
    e.save_decisions(tmp_path, AUG, review, "BPR")
    res2 = e.run(tmp_path, AUG, files, sla_days=1, bank="BPR", report_dir=tmp_path / "r")
    assert not res2.stopped_for_review and res2.review["Reviewed (Yes/No)"].iloc[0] == "Yes"


def test_a_credit_without_doc_id_is_found_only_when_irembopay_has_it(tmp_path):
    """Like TT26226M0NPS (USD 750, RFDA, 14/08): no DOC ID, but IremboPay knows the credit."""
    rows = [("2026-08-14", "TT26226M0NPS", "886000000101", "", "2026-08-14 10:13:00", 0, 750),   # 1. by reference
            ("2026-08-14", "FT20", "SZREF20", "", "2026-08-14 10:20:00", 0, 500),                # 2. desc = provider_ref
            ("2026-08-14", "FT30", "886000000103", "", "2026-08-14 10:30:00", 0, 300),           # 3. desc = invoice
            ("2026-08-14", "FT40", "886000000999", "", "2026-08-14 10:40:00", 0, 200),           # not in IremboPay
            ("2026-08-14", "FT50", "SZSHARED", "", "2026-08-14 10:50:00", 0, 100)]               # ref of two invoices
    pays = [("886000000101", 750, "USD", "2026-08-14 10:14:16", "BPR_RW", "TT26226M0NPS"),
            ("886000000102", 500, "RWF", "2026-08-14 10:20:30", "BPR_RW", "SZREF20"),
            ("886000000103", 300, "RWF", "2026-08-14 10:30:10", "BPR_RW", ""),
            ("886000000104", 100, "RWF", "2026-08-14 10:50:10", "BPR_RW", "SZSHARED"),
            ("886000000105", 100, "RWF", "2026-08-14 10:50:20", "BPR_RW", "SZSHARED")]
    res = run(tmp_path, rows=rows, pays=pays)
    assert res.to_review == 0 and not res.stopped_for_review               # found, or no invoice -> Pending
    review = res.review.copy()
    review["Reviewed (Yes/No)"] = "Yes"
    e.save_decisions(tmp_path / "base", AUG, review, "BPR")
    done = run(tmp_path, rows=rows, pays=pays)
    det = done.sheets["Payment details"].set_index("BANK_REFERENCE")
    assert det.loc["TT26226M0NPS", ["invoice_number", "Currency", "sla_status"]].tolist() == ["886000000101", "USD", "On time"]
    assert bool(det.loc["TT26226M0NPS", "notification_received"]) and bool(det.loc["FT30", "notification_received"])
    assert det.loc["FT20", bpr.FOUND_BY] == "DESCRIPTION = provider_reference" and pd.isna(det.loc["FT40", "invoice_number"])
    assert det.loc["TT26226M0NPS", bpr.FOUND_BY] == "BANK_REFERENCE = provider_reference"
    pending = set(done.sheets["Pending notification"]["BANK_REFERENCE"])
    assert {"FT40", "FT50"} <= pending and "TT26226M0NPS" not in pending      # not found -> Pending


def test_a_bpr_bank_transfer_is_bpr(tmp_path):
    # 27/10/2026: IremboPay provider 'BPR' (BANK_TRANSFER) is BPR too: received when paid, SLA checked (1 minute),
    # not in review, and the provider stays 'BPR' in Payment details
    pays = [p if p[0] != "886000000004" else ("886000000004", 300, "RWF", "2026-08-06 09:00:00", "BPR", "FT4")
            for p in PAYS]
    res = run(tmp_path, pays=pays)
    assert "FT4" not in res.review["BANK_REFERENCE"].tolist()            # no longer "paid through another channel"
    res = run(tmp_path / "b", pays=pays, require_review=False)
    t = e.counts_of(res.blocks[-1][2])
    assert t["Duplicate Payment"]["Total"] == 0 and t["Overall notifications received by Irembo"]["Total"] == 5
    d = res.sheets["Payment details"].set_index("BANK_REFERENCE")
    assert d.loc["FT4", "payment_provider"] == "BPR" and d.loc["FT4", "sla_status"] == "SLA violated"   # 3 days later
