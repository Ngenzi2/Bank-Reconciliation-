"""No transactions this period: engine.run() answers with NoTransactionsError (a ReconError) instead of a report,
and the app (app.py) turns that into the No-Transaction Report automatically - no manual choice needed, approved
2026-10-22."""
import io

import pandas as pd
import pytest
from openpyxl import load_workbook

import engine as e
from boa_fixtures import payments, statement

SEP = e.make_period("month", "2026-09-01")


def test_no_rows_in_the_period_raises_a_specific_error(tmp_path):
    # every row is in August; the period asked for is September, so none of them fall inside it
    rows = [("30/08/2026", "S1", "ePay irembo 880000000001", None, 1000)]
    files = [e.read_bank_file("s.xlsx", statement(rows), "BOA"),
            e.read_payment_file("p.csv", payments([
                dict(invoice="880000000001", amount=1000, paid_at="2026-08-30 09:00:00", ref="S1")]), "BOA")]
    with pytest.raises(e.NoTransactionsError, match="have no Irembo payments"):
        e.run(tmp_path / "base", SEP, files, require_review=False, bank="BOA", report_dir=tmp_path / "r")


def test_it_is_still_a_reconerror_for_generic_handlers():
    assert issubclass(e.NoTransactionsError, e.ReconError)


def test_no_transaction_rows_are_zero_with_the_previous_periods_count_for_comparison():
    # I&M has both currencies: both get their own row, each compared with its own count before
    previous = {"Overall processed payments": {"RWF": 10, "USD": 2, "Total": 12}}
    rows = e.no_transaction_rows("IM", previous)
    assert all(r[2] == 0 and r[3] == 0 for r in rows)                     # every category, every currency: 0 now
    rwf = next(r for r in rows if r[0] == "Overall processed payments" and r[1] == "RWF")
    assert rwf[4] == 10 and rwf[5] == "-100.00%"
    usd = next(r for r in rows if r[0] == "Overall processed payments" and r[1] == "USD")
    assert usd[4] == 2 and usd[5] == "-100.00%"


def test_no_transaction_rows_without_a_previous_report_are_na():
    rows = e.no_transaction_rows("BOA", None)
    assert all(r[4] == "N/A" and r[5] == "N/A" for r in rows)


def test_no_transaction_excel_is_a_readable_report():
    data = e.no_transaction_excel(SEP, "BOA", pd.Timestamp("2026-10-01"))
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames[0] == e.SUMMARY_SHEET["month"]
    ws = wb.active
    assert any(c.value == e.NO_TX_STATUS for row in ws.iter_rows() for c in row)
